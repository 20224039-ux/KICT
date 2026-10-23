#!/usr/bin/env python3
"""NFPA 551 8.2.1.1 (replication by an alternative calculation): the recalculated Excel calculator against qfra.py.

Reads the formula results of 외벽_QFRA_계산서.xlsx (run recalc first) and the Python point estimate, writes
results/verification.csv and report/verification.json (used by the deck).

    python check_xlsx.py
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BASE = HERE.parent


def main():
    point = pd.read_csv(BASE / "results" / "point_estimate.csv")
    wb = load_workbook(HERE / "외벽_QFRA_계산서.xlsx", data_only=True)
    ws = wb["5_위험도"]
    label = {}
    for r in range(5, ws.max_row + 1):
        b, g, name = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if b in ("B1", "B0") and g in ("M0", "M1"):
            label[(name, b, g)] = r
    rows = []
    for p in point.itertuples():
        r = label[(p.label, p.building, p.mitigation)]
        for metric, col in (("R_S", 9), ("R_P", 10), ("F_n2", 14), ("F_n1", 13), ("F_n5", 15)):
            py, xl = float(getattr(p, metric)), ws.cell(r, col).value
            if xl is None:
                raise SystemExit("no formula values: run recalc on the workbook first")
            den = max(abs(py), abs(xl))
            rows.append(dict(material=p.material, building=p.building, mitigation=p.mitigation, metric=metric,
                             python=py, excel=float(xl), rel_diff=0.0 if den == 0 else abs(py - xl) / den))
    df = pd.DataFrame(rows)
    df.to_csv(BASE / "results" / "verification.csv", index=False, encoding="utf-8")
    n_formulas = sum(1 for s in load_workbook(HERE / "외벽_QFRA_계산서.xlsx").worksheets
                     for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    out = dict(max_rel_diff=float(df.rel_diff.max()), n_cases=int(len(point)), n_values=int(len(df)), n_formulas=n_formulas)
    (HERE / "verification.json").write_text(json.dumps(out, indent=1), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
