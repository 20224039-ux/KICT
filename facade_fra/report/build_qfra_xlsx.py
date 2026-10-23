#!/usr/bin/env python3
"""Façade QFRA calculator workbook: report/외벽_QFRA_계산서.xlsx

The point estimate of qfra.py as live cell formulas, so the procedure can be followed step by step and the inputs
changed (yellow cells). Sheets follow the NFPA 551 steps:
  읽는 법 · 1_입력 (named parameters) · 2_시험자료 (t600, flag B -> tau_F, p_B) · 3_사건수목 (frequencies) ·
  4_결과모델 (storey-climb consequences per repeat test, B state, initiating event and barrier) · 5_위험도 (R = sum F C,
  exceedance, relative risk, example criterion) · 6_불확실성(MC) (Monte Carlo results of qfra.py, values) ·
  7_검증 (Python point estimate vs these formulas, NFPA 551 8.2.1.1) · 8_문서화점검 (NFPA 551 7.2, 7.3.3.4, 7.4, 8.3)

    python ../qfra.py && python build_qfra_xlsx.py
    python <xlsx skill>/scripts/recalc.py 외벽_QFRA_계산서.xlsx       # fills the formula values
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
OUT = HERE / "외벽_QFRA_계산서.xlsx"

MATERIALS = ["EPS_SC", "PIR_C", "PIR_SC", "PF_out", "PF_in"]
LABEL = {"NC": "불연 단열재(기준안)", "EPS_SC": "EPS 준불연", "PIR_C": "PIR(PUR) 가연성", "PIR_SC": "PIR(PUR) 준불연",
         "PF_out": "PF 준불연 외부용", "PF_in": "PF 준불연 내부용", "EPS_C": "EPS 가연성"}
GRADE = {"NC": "불연", "EPS_SC": "sc", "PIR_C": "c", "PIR_SC": "sc", "PF_out": "sc", "PF_in": "sc", "EPS_C": "c"}
NAME = {"N": "N_st", "h": "h_st", "units": "units", "occ": "occ", "n_aff": "n_aff", "lam1": "lam_1", "lam2": "lam_2",
        "p_spr": "p_spr", "p_fo": "p_fo", "p_vent": "p_vent", "p_x2": "p_x2", "p_exp0": "p_exp0", "p_exp1": "p_exp1",
        "p_cb": "p_cb", "p_L": "p_L", "tau_L": "tau_L_s", "t_fs": "t_fs_s", "rset": "rset_s", "s": "s_scale", "p_d": "p_d",
        "crit_f2": "crit_f2", "crit_conf": "crit_conf"}
MINUTES = {"tau_L", "t_fs", "rset", "texp1_lo", "texp1_hi", "texp2_lo", "texp2_hi"}
J_MAX = 59        # storeys above the storey of origin covered by the formulas (buildings up to 60 storeys)
N_SHOW = 20       # exceedance columns P(n >= 1..20)

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, color="4B5563")
F_INPUT = Font(name=FONT, size=10, color="0000FF")      # hardcoded inputs / measured values
F_LINK = Font(name=FONT, size=10, color="008000")       # links to another sheet
FILL_HEAD = PatternFill("solid", fgColor="23303D")
FILL_TODO = PatternFill("solid", fgColor="FFF2CC")       # cells to change
FILL_KEY = PatternFill("solid", fgColor="E2F3E2")        # key results
FILL_GREY = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SCI = "0.00E+00"
NOCLIMB = '[>=100000000]"확산 없음";0.0'       # tau_F = 1E+9 s stands for "no climb to the next storey"


def cell(ws, ref, value, font=F_BASE, fill=None, fmt=None, align=WRAP, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align
    if border:
        c.border = BOX
    return c


def rc(ws, row, col, value, **kw):
    return cell(ws, f"{get_column_letter(col)}{row}", value, **kw)


def header(ws, row, names, col0=1, widths=None):
    for k, n in enumerate(names):
        rc(ws, row, col0 + k, n, font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        if widths:
            ws.column_dimensions[get_column_letter(col0 + k)].width = widths[k]


def title(ws, text, sub=None):
    cell(ws, "A1", text, F_TITLE, border=False, align=Alignment(vertical="center"))
    if sub:
        cell(ws, "A2", sub, F_NOTE, border=False, align=Alignment(vertical="center"))


def define(wb, name, sheet, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{ref}")


def absref(col, row):
    return f"${get_column_letter(col)}${row}"


# --------------------------------------------------------------------------------------------- 읽는 법
def readme(wb):
    ws = wb.active
    ws.title = "읽는 법"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 112
    title(ws, "외벽 정량적 화재위험성평가(QFRA) 계산서 — NFPA 551 5.5 정량적 방법, 중공층 외단열 예제")
    rows = [
        ("목적", "NFPA 551(2022)이 요구하는 정량적 화재위험성평가 R = Σ F·C (식 5.5.1.2·5.5.1.3)를 외벽(중공층 외단열)에 적용하는 절차를 "
                 "칸 단위로 따라갈 수 있게 만든 계산서입니다. 노란 칸을 바꾸면 사건수목 → 결과 → 위험도 → 허용기준 판정이 다시 계산됩니다."),
        ("계산 구조", "1) 사건수목(3_사건수목): 세대 화재(IE1)·외부 발화원(IE2)이 외벽을 노출시키는 연간 빈도 F. "
                      "2) 결과 모델(4_결과모델): 중공층 시험 t₆₀₀ = '중공층 화재가 한 개 층을 올라가는 시간'으로 보고, 진압·배리어·도약 확산과 "
                      "겨루어 발화층 위로 번지는 층 수와 노출 인원 C를 계산. 3) 위험도(5_위험도): R = Σ F·C, 초과빈도 F(n), 기준안 대비 RR."),
        ("칸 색", "노란 바탕 = 바꿀 값(가정·예시값 포함) · 파란 글자 = 입력·측정값 · 검은 글자 = 계산식 · 초록 글자 = 다른 시트에서 가져온 값 · "
                  "연두 바탕 = 주요 결과"),
        ("단계 ↔ 시트", "① 문제·범위·허용기준 → 1_입력 · ② 시나리오·빈도 → 3_사건수목 · ③ 결과 → 2_시험자료, 4_결과모델 · "
                        "④ 위험도·판정 → 5_위험도 · ⑤ 불확실성 → 6_불확실성(MC) · ⑥ 검증 → 7_검증 · ⑦ 문서화·검토 → 8_문서화점검"),
        ("예시값 주의", "빈도·분기확률·시간 입력은 절차를 보이기 위한 예시값입니다(1_입력의 '근거·대체할 자료' 열). 측정값은 중공층 시험의 "
                        "t₆₀₀(3000 mm)와 플래그 B뿐입니다. 결과 수치를 그대로 인용하지 말고 국내 통계·실대형 시험값으로 바꾼 뒤 쓰십시오."),
        ("점추정과 MC", "이 시트들의 수식은 '점추정'(최적값, 중앙값, 반복시험 동일 가중)입니다. 분포를 넣은 2중 몬테카를로(인식론적 2,000 × "
                        "우연적 400 표본)는 파이썬 qfra.py가 계산하며, 그 결과는 6_불확실성(MC)에 값으로 들어 있습니다. 점추정은 꼬리 위험"
                        "(진압 지연 등)을 보지 못하므로 판정은 MC 결과로 합니다."),
        ("대책 인정 규칙", "M1(개구부 차단재 + 층간 캐비티 배리어)의 효과는 준불연 등급에만 인정합니다. 가연성 등급(PIR(PUR) 가연성)은 3 m 중공층이 "
                          "30 s 만에 전면 화염이 되었으므로, 배리어가 이를 막는다는 시험 근거가 생기기 전에는 효과를 0으로 둡니다(NFPA 551 5.1.4.5)."),
        ("제외", "EPS 가연성: 시험 기록이 4분(240·262 s)뿐이라 t₆₀₀을 구할 수 없어 정량평가에서 뺐습니다(자료 부족, 6.2.1.3). "
                 "재시험 후 2_시험자료에 넣고 '정량평가 사용'을 '예'로 바꾸면 계산에 들어갑니다(이때 4_결과모델에 행 추가 필요)."),
        ("재현", "python facade_fra/qfra.py → python facade_fra/report/build_qfra_xlsx.py → recalc. 결과 CSV: facade_fra/results/"),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        cell(ws, f"A{i}", k, F_BOLD)
        cell(ws, f"B{i}", v)
        ws.row_dimensions[i].height = 48


# --------------------------------------------------------------------------------------------- 1_입력
def inputs(wb, params):
    ws = wb.create_sheet("1_입력")
    title(ws, "① 입력값 — 건물, 빈도, 분기확률, 시간, 허용기준",
          "적용값(노란 칸)이 점추정에 쓰입니다. 하한·상한·분포는 몬테카를로(qfra.py)에서 쓰는 인식론적 불확실성 범위입니다.")
    cols = ["기호", "항목", "단위", "적용값", "하한", "상한", "분포", "계산용 값(초 단위)", "NFPA 551", "근거 · 대체할 자료"]
    header(ws, 4, cols, widths=[9, 38, 10, 12, 10, 10, 9, 14, 13, 70])
    dist_ko = {"tri": "삼각", "logu": "로그균등", "fixed": "고정"}
    groups = {"N": "건물", "lam1": "빈도", "p_spr": "분기확률", "tau_L": "시간", "s": "시험 → 실물", "p_d": "인명 환산",
              "sig_fs": "사건별 변동 (σ는 MC 전용, 지속시간 중앙값은 점추정에도 사용)", "crit_f2": "허용기준(예시)"}
    r = 5
    for p in params.itertuples():
        if p.key in groups:
            rc(ws, r, 1, groups[p.key], font=F_SUB, fill=FILL_GREY)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
            r += 1
        rc(ws, r, 1, p.symbol, font=F_BOLD, align=CENTER)
        rc(ws, r, 2, p.item)
        rc(ws, r, 3, p.unit, align=CENTER)
        sci = isinstance(p.best, float) and (abs(p.best) < 1e-2)
        fmt = SCI if sci else ("0" if float(p.best).is_integer() and p.key not in ("h",) else "0.00")
        rc(ws, r, 4, p.best, font=F_INPUT, fill=FILL_TODO, fmt=fmt, align=CENTER)
        for c, v in ((5, p.low), (6, p.high)):
            rc(ws, r, c, None if pd.isna(v) else v, font=F_INPUT, fmt=fmt, align=CENTER)
        rc(ws, r, 7, dist_ko[p.distribution], align=CENTER)
        conv = f"=D{r}*60" if p.key in MINUTES else f"=D{r}"
        rc(ws, r, 8, conv, fmt=fmt if p.key not in MINUTES else "0", align=CENTER)
        rc(ws, r, 9, p.nfpa551, align=CENTER)
        rc(ws, r, 10, p.basis)
        if p.key in NAME:
            define(wb, NAME[p.key], "1_입력", absref(8, r))
        if p.key in ("texp1_lo", "texp1_hi", "texp2_lo", "texp2_hi"):
            define(wb, p.key + "_s", "1_입력", absref(8, r))
        r += 1
    r += 1
    rc(ws, r, 1, "파생값", font=F_SUB, fill=FILL_GREY)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    r += 1
    derived = [
        ("t_e1", "분출화염 지속(IE1, 중앙값) — 비자립 재료의 확산 종료 시각", "s", "=(texp1_lo_s+texp1_hi_s)/2", "texp1_s", "0"),
        ("t_e2", "외부 화원 지속(IE2, 중앙값)", "s", "=(texp2_lo_s+texp2_hi_s)/2", "texp2_s", "0"),
        ("o_st", "확산 띠 안의 층당 재실자 = n_a × o", "명", "=n_aff*occ", "occ_st", "0.0"),
        ("O", "건물 재실자 = N × n_u × o", "명", "=N_st*units*occ", "occ_total", "0"),
        ("U", "세대 수 = N × n_u", "세대", "=N_st*units", "n_units", "0"),
    ]
    for sym, item, unit, f, name, fmt in derived:
        rc(ws, r, 1, sym, font=F_BOLD, align=CENTER)
        rc(ws, r, 2, item)
        rc(ws, r, 3, unit, align=CENTER)
        rc(ws, r, 8, f, fmt=fmt, align=CENTER)
        define(wb, name, "1_입력", absref(8, r))
        r += 1
    ws.freeze_panes = "A5"
    return ws


# --------------------------------------------------------------------------------------------- 2_시험자료
def tests_sheet(wb, tests):
    ws = wb.create_sheet("2_시험자료")
    title(ws, "③-1 시험 자료 → 결과 모델 입력 (중공층 100 mm · 외단열 100 mm · 30 kW 20분 · 30분 평가)",
          "t₆₀₀ = 3000 mm에서 ΔT ≥ 600 K가 시작되는 시각 = 중공층 화재가 한 개 층(3 m)을 올라가는 시간(FRI 개념). "
          "플래그 B = 버너 소화 후 자립 연소(6개 높이 중 3곳 이상 ssr ≥ 1).")
    cols = ["시험", "재료", "등급", "t₆₀₀ 3000 mm (s)", "플래그 B (1/0)", "정량평가 사용", "τ_F = s·t₆₀₀·h/3 m (s)", "비고"]
    header(ws, 4, cols, widths=[11, 18, 7, 14, 12, 11, 17, 62])
    rows = {}
    for i, t in enumerate(tests.itertuples(), start=5):
        rows[t.test] = i
        rc(ws, i, 1, t.test, font=F_BOLD)
        rc(ws, i, 2, t.label)
        rc(ws, i, 3, t.grade, align=CENTER)
        rc(ws, i, 4, None if pd.isna(t.t600_s) else float(t.t600_s), font=F_INPUT, fmt="0", align=CENTER)
        rc(ws, i, 5, int(t.flag_B), font=F_INPUT, align=CENTER)
        rc(ws, i, 6, "예" if t.used else "아니오", font=F_INPUT, align=CENTER)
        rc(ws, i, 7, f'=IF(OR(F{i}="아니오",D{i}=""),1E+9,s_scale*D{i}*h_st/3)', fmt=NOCLIMB, align=CENTER)
        rc(ws, i, 8, t.note if isinstance(t.note, str) else "")
    last = 4 + len(tests)
    note_r = last + 1
    cell(ws, f"A{note_r}", "τ_F = 1E+9 s: 층간 확산 없음(600 K 미도달 또는 평가 제외). 정량평가 사용이 '아니오'인 시험은 4_결과모델에 들어가지 않습니다.",
         F_NOTE, border=False, align=Alignment(vertical="center"))
    r0 = last + 3
    cell(ws, f"A{r0 - 1}", "재료별 요약 — p_B는 Jeffreys 추정 (k + ½)/(n + 1): 반복 2–3회의 0/2, 2/2를 0·1로 단정하지 않기 위함",
         F_SUB, border=False, align=Alignment(vertical="center"))
    header(ws, r0, ["재료", "재료명", "등급", "시험 수 n", "B 횟수 k", "p_B", "t₆₀₀ 평균 (s)", "해석"])
    summary = {}
    interp = {"EPS_SC": "용융형: 3회 중 1회는 30분 안에 600 K 미도달, 자립 연소 없음",
              "PIR_C": "가연성: 30 s 만에 전면 화염 → 층당 시간이 가장 짧음",
              "PIR_SC": "탄화형, 자립 연소 없음",
              "PF_out": "탄화형, 2회 모두 자립 연소(B) → 진압이 늦으면 계속 올라감",
              "PF_in": "탄화형, 확산 빠름(III등급), 2회 중 1회 자립 연소"}
    rng = f"$B$5:$B${last}"
    for k, m in enumerate(MATERIALS):
        r = r0 + 1 + k
        summary[m] = r
        rc(ws, r, 1, m, font=F_BOLD)
        rc(ws, r, 2, LABEL[m])
        rc(ws, r, 3, GRADE[m], align=CENTER)
        rc(ws, r, 4, f'=COUNTIFS({rng},B{r},$F$5:$F${last},"예")', align=CENTER)
        rc(ws, r, 5, f'=SUMIFS($E$5:$E${last},{rng},B{r},$F$5:$F${last},"예")', align=CENTER)
        rc(ws, r, 6, f"=(E{r}+0.5)/(D{r}+1)", fmt="0.000", align=CENTER, fill=FILL_KEY)
        rc(ws, r, 7, f'=IFERROR(AVERAGEIFS($D$5:$D${last},{rng},B{r},$F$5:$F${last},"예"),"")', fmt="0", align=CENTER)
        rc(ws, r, 8, interp[m])
    return rows, summary


# --------------------------------------------------------------------------------------------- 3_사건수목
def event_tree(wb):
    ws = wb.create_sheet("3_사건수목")
    title(ws, "② 사건수목 — 외벽이 화염에 노출되는 연간 빈도 (동·년 기준)",
          "분기확률은 서로 독립이라고 보고 경로를 따라 곱합니다(NFPA 551 5.3.1.2.2). B1 = 세대 스프링클러 설치, B0 = 미설치.")
    cols = ["단계", "사건 / 분기", "분기확률 (식)", "B1 (/년)", "B0 (/년)", "설명"]
    header(ws, 4, cols, widths=[7, 40, 28, 14, 14, 60])
    rows = [
        ("IE1", "세대 화재 발생", "λ₁ × 세대 수", "=lam_1*n_units", "=lam_1*n_units", "초기사건(IE1): 실내 발화"),
        ("", "스프링클러 제어 실패(또는 미설치)", "1 − p_s (B0: 1)", "=D5*(1-p_spr)", "=E5", "제어되면 실내 국한(외벽 영향 없음)"),
        ("", "최성기(플래시오버) 도달", "p_FO", "=D6*p_fo", "=E6*p_fo", ""),
        ("", "창 파손 · 외부 분출화염 = F₁", "p_v", "=D7*p_vent", "=E7*p_vent", "외벽 노출 사건(IE1)의 빈도 F₁"),
        ("", "   ├ 중공층·단열재 노출 (M0)", "p_e0", "=D8*p_exp0", "=E8*p_exp0", "→ 외벽 경로 + 창→창 도약 (4_결과모델: IE1)"),
        ("", "   ├ 노출 안 됨 (M0)", "1 − p_e0", "=D8*(1-p_exp0)", "=E8*(1-p_exp0)", "→ 창→창 도약만"),
        ("", "   ├ 중공층·단열재 노출 (M1, 준불연만)", "p_e1", "=D8*p_exp1", "=E8*p_exp1", "개구부 불연 차단재가 있을 때"),
        ("", "   └ 노출 안 됨 (M1, 준불연만)", "1 − p_e1", "=D8*(1-p_exp1)", "=E8*(1-p_exp1)", ""),
        ("IE2", "외부 발화원 화재(외벽 하부)", "λ₂", "=lam_2", "=lam_2", "초기사건(IE2): 필로티·주차 차량, 쓰레기, 외부 설비"),
        ("", "외벽 화염 노출 = F₂", "p_x", "=D13*p_x2", "=E13*p_x2", "외벽 노출 사건(IE2)의 빈도 F₂ (스프링클러 무관)"),
        ("", "   ├ 중공층·단열재 노출 (M0)", "p_e0", "=D14*p_exp0", "=E14*p_exp0", "→ 외벽 경로만 (도약 없음; 4_결과모델: IE2)"),
        ("", "   ├ 노출 안 됨 (M0)", "1 − p_e0", "=D14*(1-p_exp0)", "=E14*(1-p_exp0)", "→ 확산 없음"),
        ("", "   ├ 중공층·단열재 노출 (M1, 준불연만)", "p_e1", "=D14*p_exp1", "=E14*p_exp1", ""),
        ("", "   └ 노출 안 됨 (M1, 준불연만)", "1 − p_e1", "=D14*(1-p_exp1)", "=E14*(1-p_exp1)", ""),
    ]
    for i, (a, b, c, d, e, f) in enumerate(rows, start=5):
        rc(ws, i, 1, a, font=F_BOLD, align=CENTER)
        rc(ws, i, 2, b, font=F_BOLD if b.startswith(("창", "외벽 화염")) else F_BASE)
        rc(ws, i, 3, c, align=CENTER)
        key = i in (8, 14)
        rc(ws, i, 4, d, fmt=SCI, align=CENTER, fill=FILL_KEY if key else None)
        rc(ws, i, 5, e, fmt=SCI, align=CENTER, fill=FILL_KEY if key else None)
        rc(ws, i, 6, f)
    define(wb, "F_IE1_B1", "3_사건수목", "$D$8")
    define(wb, "F_IE1_B0", "3_사건수목", "$E$8")
    define(wb, "F_IE2", "3_사건수목", "$D$14")
    cell(ws, "A20", "불연 단열재(기준안)는 중공층 경로가 없으므로 IE1은 모두 '도약만', IE2는 확산 없음으로 계산합니다. "
                    "가연성 등급은 M1이어도 M0의 분기확률(p_e0, 배리어 없음)을 씁니다.", F_NOTE, border=False,
         align=Alignment(vertical="center"))


# --------------------------------------------------------------------------------------------- 4_결과모델
def consequences(wb, tests, test_rows, summary):
    ws = wb.create_sheet("4_결과모델")
    title(ws, "③-2 결과 모델 — 발화층 위로 번지는 층 수와 노출 인원 (사건 1회당 기대값)",
          "n_F = ⌊min(T_FS, T_stop)/τ_F⌋, a_j = q^j·[j ≤ n_F], b_j = p_L^j·[j ≤ n_L], P(n ≥ j) = w_j·(1 − (1 − a_j)(1 − b_j)), "
          "w_j = (N − j)/N (IE1: 발화층 균등) 또는 1 (IE2: 1층 발화). 층 j는 먼저 닿는 경로의 도달시각이 RSET보다 이르면 노출.")
    # helper rows: j, weights
    rc(ws, 4, 1, "j (발화층 위 층)", font=F_BOLD)
    rc(ws, 5, 1, "w_j (IE1)", font=F_BOLD)
    rc(ws, 6, 1, "w_j (IE2)", font=F_BOLD)
    for j in range(1, J_MAX + 1):
        c = 1 + j
        rc(ws, 4, c, j, font=F_INPUT, align=CENTER)
        rc(ws, 5, c, f"=IF({get_column_letter(c)}4<N_st,(N_st-{get_column_letter(c)}4)/N_st,0)", fmt="0.000", align=CENTER)
        rc(ws, 6, c, f"=IF({get_column_letter(c)}4<N_st,1,0)", align=CENTER)
    last_col = get_column_letter(1 + J_MAX)
    define(wb, "j_row", "4_결과모델", f"$B$4:${last_col}$4")
    define(wb, "w_ie1", "4_결과모델", f"$B$5:${last_col}$5")
    define(wb, "w_ie2", "4_결과모델", f"$B$6:${last_col}$6")
    for c in range(2, 2 + J_MAX):
        ws.column_dimensions[get_column_letter(c)].width = 9

    # event rows
    h0 = 9
    cols = ["재료", "시험", "사건", "배리어(1/0)", "B 상태(1/0)", "가중치", "τ_F (s)", "T_stop (s)", "n_F", "n_L",
            "q = 1 − p_cb", "p_L(사건)", "E[확산 층 수]", "E[노출 인원]"]
    cell(ws, f"A{h0 - 1}", "사건 행: 반복시험 × B 상태 × 사건(IE1/IE2) × 배리어. 가중치 = (1/n) × (B면 p_B, 아니면 1 − p_B). "
                           "'L'은 중공층 경로가 없는 창→창 도약만의 경로(IE1, 불연 기준안과 '노출 안 됨' 분기에 사용).",
         F_NOTE, border=False, align=Alignment(vertical="center"))
    header(ws, h0, cols)
    for k in range(N_SHOW):
        rc(ws, h0 - 1, 15 + k, k + 1, font=F_INPUT, align=CENTER)     # n for the exceedance columns
        rc(ws, h0, 15 + k, f"P(n≥{k + 1})", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
    widths = [10, 10, 6, 8, 8, 9, 10, 10, 6, 6, 9, 9, 12, 12]
    for k, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(k)].width = max(wdt, 9)

    events = []
    for m in MATERIALS:
        for t in tests[(tests.material == m) & tests.used].itertuples():
            for ie in (1, 2):
                for bar in (0, 1):
                    for b in (1, 0):
                        events.append((m, t.test, ie, bar, b))
    events.append(("L", "—", 1, 0, 0))
    r = h0 + 1
    first = r
    for m, test, ie, bar, b in events:
        rc(ws, r, 1, m, font=F_BOLD)
        rc(ws, r, 2, test, align=CENTER)
        rc(ws, r, 3, ie, align=CENTER)
        rc(ws, r, 4, bar, align=CENTER)
        rc(ws, r, 5, b, align=CENTER)
        if m == "L":
            rc(ws, r, 6, 1, align=CENTER)
            rc(ws, r, 7, "=1E+9", fmt=NOCLIMB, align=CENTER)
        else:
            sr = summary[m]
            n_ref, pb_ref = f"'2_시험자료'!$D${sr}", f"'2_시험자료'!$F${sr}"
            pb = pb_ref if b == 1 else f"(1-{pb_ref})"
            rc(ws, r, 6, f"=(1/{n_ref})*{pb}", font=F_LINK, fmt="0.000", align=CENTER)
            rc(ws, r, 7, f"='2_시험자료'!G{test_rows[test]}", font=F_LINK, fmt=NOCLIMB, align=CENTER)
        texp = "texp1_s" if ie == 1 else "texp2_s"
        rc(ws, r, 8, f"=IF(E{r}=1,1E+9,{texp})", fmt='[>=100000000]"자립(∞)";0', align=CENTER)
        rc(ws, r, 9, f"=FLOOR(MIN(t_fs_s,H{r})/G{r},1)", align=CENTER)
        rc(ws, r, 10, f"=IF(C{r}=1,FLOOR(t_fs_s/tau_L_s,1),0)", align=CENTER)
        rc(ws, r, 11, f"=IF(D{r}=1,1-p_cb,1)", fmt="0.00", align=CENTER)
        rc(ws, r, 12, f"=IF(C{r}=1,p_L,0)", fmt="0.00", align=CENTER)
        w = "w_ie1" if ie == 1 else "w_ie2"
        a = f"(j_row<=I{r})*K{r}^j_row"
        bb = f"(j_row<=J{r})*L{r}^j_row"
        ta, tb = f"(j_row*G{r}<rset_s)", "(j_row*tau_L_s<rset_s)"
        rc(ws, r, 13, f"=SUMPRODUCT({w},1-(1-{a})*(1-{bb}))", fmt="0.0000", align=CENTER)
        seen = f"{a}*{bb}*(({ta}+{tb})>0)+{a}*(1-{bb})*{ta}+(1-{a})*{bb}*{tb}"
        rc(ws, r, 14, f"=occ_st*SUMPRODUCT({w},{seen})", fmt="0.000", align=CENTER)
        for k in range(N_SHOW):
            col = 15 + k
            nref = f"{get_column_letter(col)}${h0 - 1}"
            wn = f"IF({nref}<N_st,(N_st-{nref})/N_st,0)" if ie == 1 else f"IF({nref}<N_st,1,0)"
            rc(ws, r, col, f"={wn}*(1-(1-({nref}<=I{r})*K{r}^{nref})*(1-({nref}<=J{r})*L{r}^{nref}))", fmt="0.0000", align=CENTER)
        r += 1
    last = r - 1

    # aggregation: per material x IE x barrier (and the leapfrog-only path)
    a0 = last + 3
    cell(ws, f"A{a0 - 1}", "재료별 집계 = Σ 가중치 × 사건 행 (5_위험도가 여기서 가져갑니다)", F_SUB, border=False,
         align=Alignment(vertical="center"))
    header(ws, a0, ["재료", "사건", "배리어", "E[확산 층 수]", "E[노출 인원]"] + [f"P(n≥{k + 1})" for k in range(N_SHOW)])
    rows = [(m, ie, bar) for m in MATERIALS for ie in (1, 2) for bar in (0, 1)] + [("L", 1, 0)]
    rr = a0 + 1
    for m, ie, bar in rows:
        rc(ws, rr, 1, m, font=F_BOLD)
        rc(ws, rr, 2, ie, align=CENTER)
        rc(ws, rr, 3, bar, align=CENTER)
        mask = f"($A${first}:$A${last}=A{rr})*($C${first}:$C${last}=B{rr})*($D${first}:$D${last}=C{rr})*$F${first}:$F${last}"
        for k, src in enumerate(["M", "N"] + [get_column_letter(15 + q) for q in range(N_SHOW)]):
            fmt = "0.0000" if k != 1 else "0.000"
            rc(ws, rr, 4 + k, f"=SUMPRODUCT({mask}*{src}${first}:{src}${last})", fmt=fmt, align=CENTER,
               fill=FILL_KEY if k < 2 else None)
        rr += 1
    agg_first, agg_last = a0 + 1, rr - 1
    define(wb, "agg_mat", "4_결과모델", f"$A${agg_first}:$A${agg_last}")
    define(wb, "agg_ie", "4_결과모델", f"$B${agg_first}:$B${agg_last}")
    define(wb, "agg_bar", "4_결과모델", f"$C${agg_first}:$C${agg_last}")
    for k, nm in enumerate(["agg_En", "agg_Ep"] + [f"agg_P{q + 1}" for q in range(N_SHOW)]):
        col = get_column_letter(4 + k)
        define(wb, nm, "4_결과모델", f"${col}${agg_first}:${col}${agg_last}")
    ws.freeze_panes = f"C{h0 + 1}"


def lk(mat, ie, bar, col):
    """value of the aggregation block for (material, IE, barrier) -- 0 when absent."""
    return f"SUMPRODUCT((agg_mat={mat})*(agg_ie={ie})*(agg_bar={bar})*{col})"


# --------------------------------------------------------------------------------------------- 5_위험도
def risk_sheet(wb, point):
    ws = wb.create_sheet("5_위험도")
    title(ws, "④ 위험도 R = Σ F·C (NFPA 551 식 5.5.1.3), 초과빈도, 기준안 대비 상대위험도, 허용기준 판정(예시)",
          "R_S = F₁·[p_e·C(IE1 외벽+도약) + (1 − p_e)·C(도약만)] + F₂·p_e·C(IE2 외벽). C는 4_결과모델 집계값. "
          "판정은 점추정이며, 불확실성을 넣은 판정은 6_불확실성(MC)를 보십시오.")
    cols = ["건물", "대책", "재료", "등급", "대책 인정(1/0)", "F₁ 분출 (/년)", "F₂ 외부 노출 (/년)", "p_e",
            "R_S 확산 층 수 (/년)", "R_P 노출 인원 (/년)", "R_fatal 사망 (/년)", "개인위험도 IR (/년)",
            "F(n≥1) (/년)", "F(n≥2) (/년)", "F(n≥5) (/년)", "RR_S (기준안 대비)", "IE2 비중", "허용기준 예시 F(n≥2) ≤ F₂*"]
    header(ws, 4, cols, widths=[7, 7, 18, 7, 9, 12, 12, 8, 13, 13, 13, 13, 12, 12, 12, 11, 9, 14])
    r = 5
    rowmap = {}
    for b in ("B1", "B0"):
        for g in ("M0", "M1"):
            for m in ["NC"] + MATERIALS:
                rowmap[(m, b, g)] = r
                rc(ws, r, 1, b, align=CENTER)
                rc(ws, r, 2, g, align=CENTER)
                rc(ws, r, 3, LABEL[m], font=F_BOLD)
                rc(ws, r, 4, GRADE[m], align=CENTER)
                rc(ws, r, 5, f'=IF(AND(B{r}="M1",D{r}="sc"),1,0)', align=CENTER)
                rc(ws, r, 6, "=F_IE1_B1" if b == "B1" else "=F_IE1_B0", font=F_LINK, fmt=SCI, align=CENTER)
                rc(ws, r, 7, "=F_IE2", font=F_LINK, fmt=SCI, align=CENTER)
                rc(ws, r, 8, f'=IF(D{r}="불연",0,IF(E{r}=1,p_exp1,p_exp0))', fmt="0.00", align=CENTER)
                mat = f'"{m}"'
                def risk_of(col):
                    if m == "NC":
                        return f"F{r}*{lk(chr(34) + 'L' + chr(34), 1, 0, col)}"
                    c1 = lk(mat, 1, f"E{r}", col)
                    cl = lk('"L"', 1, 0, col)
                    c2 = lk(mat, 2, f"E{r}", col)
                    return f"F{r}*(H{r}*{c1}+(1-H{r})*{cl})+G{r}*H{r}*{c2}"
                rc(ws, r, 9, "=" + risk_of("agg_En"), fmt=SCI, align=CENTER, fill=FILL_KEY)
                rc(ws, r, 10, "=" + risk_of("agg_Ep"), fmt=SCI, align=CENTER, fill=FILL_KEY)
                rc(ws, r, 11, f"=p_d*J{r}", fmt=SCI, align=CENTER)
                rc(ws, r, 12, f"=K{r}/occ_total", fmt=SCI, align=CENTER)
                rc(ws, r, 13, "=" + risk_of("agg_P1"), fmt=SCI, align=CENTER)
                rc(ws, r, 14, "=" + risk_of("agg_P2"), fmt=SCI, align=CENTER, fill=FILL_KEY)
                rc(ws, r, 15, "=" + risk_of("agg_P5"), fmt=SCI, align=CENTER)
                r += 1
    # relative risk and IE2 share need the baseline row of the same block
    for (m, b, g), rr in rowmap.items():
        base = rowmap[("NC", b, g)]
        rc(ws, rr, 16, f"=I{rr}/I${base}", fmt="0.0", align=CENTER)
        if m == "NC":
            rc(ws, rr, 17, 0, fmt="0%", align=CENTER)
        else:
            c2 = lk(f'"{m}"', 2, f"E{rr}", "agg_En")
            rc(ws, rr, 17, f"=G{rr}*H{rr}*{c2}/I{rr}", fmt="0%", align=CENTER)
        rc(ws, rr, 18, f'=IF(N{rr}<=crit_f2,"충족","불충족")', align=CENTER)
    last = r - 1
    cell(ws, f"A{last + 2}", "R_S: 외벽 때문에 발화층 위로 번지는 층 수의 연간 기대값 · R_P: 그 층들에서 피난 전에 화염이 닿는 재실자 수의 연간 기대값 · "
                             "R_fatal = p_d × R_P · IR = R_fatal ÷ 건물 재실자 (참고: 영국 HSE R2P2 개인위험도 10⁻⁶/년 '일반적으로 허용', "
                             "10⁻⁴/년 '허용 한계').", F_NOTE, border=False, align=Alignment(vertical="center"))

    # exceedance table (B1) for M0 and M1, n = 1 .. N_SHOW
    e0 = last + 5
    cell(ws, f"A{e0 - 1}", "초과빈도 F(발화층 위 확산 ≥ n) — 스프링클러 설치 건물(B1), 점추정", F_SUB, border=False,
         align=Alignment(vertical="center"))
    header(ws, e0, ["대책", "", "재료"] + [f"n = {k + 1}" for k in range(N_SHOW)])
    er = e0 + 1
    for g in ("M0", "M1"):
        for m in ["NC"] + MATERIALS:
            src = rowmap[(m, "B1", g)]
            rc(ws, er, 1, g, align=CENTER)
            rc(ws, er, 2, "", align=CENTER)
            rc(ws, er, 3, LABEL[m], font=F_BOLD)
            for k in range(N_SHOW):
                col = f"agg_P{k + 1}"
                if m == "NC":
                    f = f"=$F${src}*{lk(chr(34) + 'L' + chr(34), 1, 0, col)}"
                else:
                    c1 = lk(f'"{m}"', 1, f"$E${src}", col)
                    cl = lk('"L"', 1, 0, col)
                    c2 = lk(f'"{m}"', 2, f"$E${src}", col)
                    f = f"=$F${src}*($H${src}*{c1}+(1-$H${src})*{cl})+$G${src}*$H${src}*{c2}"
                rc(ws, er, 4 + k, f, fmt=SCI, align=CENTER)
            er += 1
    for k in range(N_SHOW):
        ws.column_dimensions[get_column_letter(4 + k)].width = max(ws.column_dimensions[get_column_letter(4 + k)].width or 0, 11)

    # chart: R_S, B1, M0 vs M1 (log scale)
    c0, cc = 3, 20                                 # chart table at T3, right of the main table
    rc(ws, c0, cc, "그림용 표: R_S (B1)", font=F_SUB, border=False)
    header(ws, c0 + 1, ["재료", "M0 대책 없음", "M1 차단재 + 배리어"], col0=cc, widths=[20, 13, 15])
    for k, m in enumerate(["NC"] + MATERIALS):
        rr = c0 + 2 + k
        rc(ws, rr, cc, LABEL[m])
        rc(ws, rr, cc + 1, f"=I{rowmap[(m, 'B1', 'M0')]}", fmt=SCI, align=CENTER)
        rc(ws, rr, cc + 2, f"=I{rowmap[(m, 'B1', 'M1')]}", fmt=SCI, align=CENTER)
    ch = BarChart()
    ch.type = "col"
    ch.title = "외벽 확산 위험도 R_S (층/동·년, 점추정, B1)"
    ch.y_axis.title = "R_S (log)"
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.scaling.min = 1e-6
    ch.y_axis.majorGridlines = None
    data = Reference(ws, min_col=cc + 1, max_col=cc + 2, min_row=c0 + 1, max_row=c0 + 7)
    cats = Reference(ws, min_col=cc, min_row=c0 + 2, max_row=c0 + 7)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.height, ch.width = 8, 18
    for s, colr in zip(ch.series, ("23303D", "9CA3AF")):
        s.graphicalProperties.solidFill = colr
        s.graphicalProperties.line.solidFill = colr
    ws.add_chart(ch, f"{get_column_letter(cc)}{c0 + 9}")
    ws.freeze_panes = "D5"
    return rowmap


# --------------------------------------------------------------------------------------------- 6_불확실성(MC)
def mc_sheet(wb, mc, ex, sens, params):
    ws = wb.create_sheet("6_불확실성(MC)")
    title(ws, "⑤ 불확실성 — 2중 몬테카를로 결과 (qfra.py, 값)",
          "바깥 루프 = 인식론적 불확실성(모르는 값: 빈도·분기확률·s·중앙값·p_B) 2,000개, 안쪽 루프 = 우연적 변동(사건마다 다른 값: 진압·피난 시각, "
          "도약 시간, 분출 지속, 반복시험, B) 400개(라틴 하이퍼큐브). 입력을 바꾸면 qfra.py를 다시 실행해야 이 시트가 바뀝니다.")
    cols = ["건물", "대책", "재료", "R_S 평균", "R_S 5 %", "R_S 50 %", "R_S 95 %", "R_P 평균", "F(n≥2) 평균", "F(n≥2) 50 %",
            "F(n≥2) 95 %", "RR_S 50 %", "RR_S 5–95 %", "기준 충족 비율", "판정 (신뢰수준 90 %)"]
    header(ws, 4, cols, widths=[7, 7, 18, 11, 11, 11, 11, 11, 12, 12, 12, 10, 16, 12, 16])
    r = 5
    for row in mc.itertuples():
        rc(ws, r, 1, row.building, align=CENTER)
        rc(ws, r, 2, row.mitigation, align=CENTER)
        rc(ws, r, 3, row.label, font=F_BOLD)
        for k, v in enumerate([row.R_S_mean, row.R_S_p05, row.R_S_p50, row.R_S_p95, row.R_P_mean, row.F_n2_mean,
                               row.F_n2_p50, row.F_n2_p95]):
            rc(ws, r, 4 + k, float(v), font=F_INPUT, fmt=SCI, align=CENTER)
        rc(ws, r, 12, float(row.RR_S_p50), font=F_INPUT, fmt="0.0", align=CENTER)
        rc(ws, r, 13, f"{row.RR_S_p05:.1f} – {row.RR_S_p95:.0f}", font=F_INPUT, align=CENTER)
        rc(ws, r, 14, float(row.conf_meets_F2), font=F_INPUT, fmt="0.0%", align=CENTER, fill=FILL_KEY)
        rc(ws, r, 15, row.verdict, font=F_BOLD, align=CENTER, fill=FILL_KEY)
        r += 1
    cell(ws, f"A{r}", "판정: 인식론적 표본 중 F(n≥2) ≤ F₂*를 만족하는 비율이 90 % 이상이면 '충족', 10 % 이하면 '불충족', 그 사이는 '판단 유보'"
                      "(자료를 더 모아야 결론이 나는 상태). 기준값과 신뢰수준은 예시입니다.", F_NOTE, border=False,
         align=Alignment(vertical="center"))

    # sensitivity: top parameters per material (B1)
    s0 = r + 3
    cell(ws, f"A{s0 - 1}", "민감도 — log R_S와 인식론적 입력의 순위상관 ρ (B1). 절댓값이 큰 입력부터 자료를 보강하면 판정이 가장 빨리 확정됩니다.",
         F_SUB, border=False, align=Alignment(vertical="center"))
    sym = {p.key: p.symbol for p in params.itertuples()}
    sym["p_B"] = "p_B"
    item = {p.key: p.item for p in params.itertuples()}
    item["p_B"] = "자립 연소 확률(플래그 B)"
    keys = [k for k in sens.parameter.unique() if k != "p_d"]
    header(ws, s0, ["대책", "입력", "항목"] + [LABEL[m] for m in MATERIALS])
    rr = s0 + 1
    for g in ("M0", "M1"):
        sub = sens[sens.mitigation == g]
        order = sub.groupby("parameter").rho_R_S.apply(lambda v: v.abs().max()).sort_values(ascending=False).index
        for k in [x for x in order if x in keys]:
            rc(ws, rr, 1, g, align=CENTER)
            rc(ws, rr, 2, sym[k], font=F_BOLD, align=CENTER)
            rc(ws, rr, 3, item[k])
            for q, m in enumerate(MATERIALS):
                v = float(sub[(sub.material == m) & (sub.parameter == k)].rho_R_S.iloc[0])
                rc(ws, rr, 4 + q, v, font=F_INPUT, fmt="+0.00;-0.00;0.00", align=CENTER,
                   fill=FILL_TODO if abs(v) >= 0.3 else None)
            rr += 1
    cell(ws, f"A{rr}", "노란 바탕: |ρ| ≥ 0.3 (주요 불확실성).", F_NOTE, border=False, align=Alignment(vertical="center"))

    # exceedance mean curves (B1, M0) and chart
    e0 = rr + 3
    cell(ws, f"A{e0 - 1}", "초과빈도 평균 곡선 F(n) — B1, M0 (MC 평균)", F_SUB, border=False, align=Alignment(vertical="center"))
    header(ws, e0, ["n"] + [LABEL[m] for m in ["NC"] + MATERIALS])
    ns = sorted(ex.n.unique())
    for k, n in enumerate(ns):
        rc(ws, e0 + 1 + k, 1, int(n), align=CENTER)
        for q, m in enumerate(["NC"] + MATERIALS):
            v = ex[(ex.material == m) & (ex.building == "B1") & (ex.mitigation == "M0") & (ex.n == n)].F_mean.iloc[0]
            rc(ws, e0 + 1 + k, 2 + q, max(float(v), 1e-12), font=F_INPUT, fmt=SCI, align=CENTER)
    sc = ScatterChart()
    sc.title = "초과빈도 F(확산 층 수 ≥ n), B1·M0, MC 평균"
    sc.style = 13
    sc.x_axis.title = "발화층 위 확산 층 수 n"
    sc.y_axis.title = "F (/동·년, log)"
    sc.y_axis.scaling.logBase = 10
    sc.y_axis.scaling.min = 1e-8
    sc.y_axis.scaling.max = 1e-2
    xref = Reference(ws, min_col=1, min_row=e0 + 1, max_row=e0 + len(ns))
    colors = ["8A8F98", "E87BA4", "6250D6", "008300", "2A78D6", "E34948"]
    for q, colr in enumerate(colors):
        yref = Reference(ws, min_col=2 + q, min_row=e0, max_row=e0 + len(ns))
        s = Series(yref, xref, title_from_data=True)
        s.graphicalProperties.line.solidFill = colr
        s.graphicalProperties.line.width = 22000
        s.marker.symbol = "none"
        sc.series.append(s)
    sc.height, sc.width = 9, 18
    ws.add_chart(sc, f"I{e0}")


# --------------------------------------------------------------------------------------------- 7_검증
def verify_sheet(wb, point, rowmap):
    ws = wb.create_sheet("7_검증")
    title(ws, "⑥ 검증 — 대체 계산에 의한 재현 (NFPA 551 8.2.1.1)",
          "파이썬 qfra.py의 점추정(값)과 이 계산서의 수식 결과를 비교합니다. 두 구현이 독립적으로 같은 값을 내면 계산 단계의 오류가 없다는 근거가 됩니다. "
          "1_입력을 바꾸면 파이썬 값과 달라지는 것이 정상입니다(그때는 qfra.py의 PARAMS도 같이 바꾸고 다시 실행).")
    cols = ["건물", "대책", "재료", "R_S 파이썬", "R_S 계산서", "상대차", "R_P 파이썬", "R_P 계산서", "상대차",
            "F(n≥2) 파이썬", "F(n≥2) 계산서", "상대차"]
    header(ws, 4, cols, widths=[7, 7, 18, 13, 13, 10, 13, 13, 10, 13, 13, 10])
    r = 5
    for p in point.itertuples():
        src = rowmap[(p.material, p.building, p.mitigation)]
        rc(ws, r, 1, p.building, align=CENTER)
        rc(ws, r, 2, p.mitigation, align=CENTER)
        rc(ws, r, 3, p.label, font=F_BOLD)
        for k, (val, col) in enumerate(((p.R_S, "I"), (p.R_P, "J"), (p.F_n2, "N"))):
            c = 4 + 3 * k
            rc(ws, r, c, float(val), font=F_INPUT, fmt=SCI, align=CENTER)
            rc(ws, r, c + 1, f"='5_위험도'!{col}{src}", font=F_LINK, fmt=SCI, align=CENTER)
            a, b = get_column_letter(c), get_column_letter(c + 1)
            rc(ws, r, c + 2, f"=IF(MAX(ABS({a}{r}),ABS({b}{r}))=0,0,ABS({a}{r}-{b}{r})/MAX(ABS({a}{r}),ABS({b}{r})))",
               fmt="0.0E+00", align=CENTER)
        r += 1
    last = r - 1
    rc(ws, r + 1, 3, "최대 상대차", font=F_BOLD)
    rc(ws, r + 1, 4, f"=MAX(F5:F{last},I5:I{last},L5:L{last})", fmt="0.0E+00", align=CENTER, fill=FILL_KEY)
    rc(ws, r + 1, 5, f'=IF(D{r + 1}<1E-9,"일치","불일치 — 입력 확인")', font=F_BOLD, align=CENTER, fill=FILL_KEY)


# --------------------------------------------------------------------------------------------- 8_문서화점검
CHECK = [
    ("개념보고서", "7.2.2.1", "참여자와 자격(학력·FRA 경력·자격) 목록", "보고서 1.1"),
    ("개념보고서", "7.2.2.2", "범위와 목적: 평가 목적, 대상 건물·외벽 시스템, 평가 한계(사용 단계, 외벽 경로만) — A.7.2.2.2", "보고서 1.2 · 1_입력(건물)"),
    ("개념보고서", "7.2.2.3", "목표: 인명(발화층 위 재실자), 소방대, 재산 중 무엇을 지키는가", "보고서 1.3"),
    ("개념보고서", "7.2.2.4", "허용기준: 지표·기준값·신뢰수준(사전협의 결과)", "1_입력(허용기준) · 보고서 1.4"),
    ("개념보고서", "7.2.2.5", "위험요소: 노출 대상·화재 자극·전달 현상·대상의 반응(4.4.2)", "보고서 1.5"),
    ("개념보고서", "7.2.2.6", "시나리오와 군집: IE1·IE2, 군집 근거, 제외한 시나리오와 이유", "3_사건수목 · 보고서 1.6"),
    ("개념보고서", "7.2.2.7", "방법: 사건수목 + 층간 확산 결과 모델 + 2중 MC, 적합성 근거", "보고서 1.7"),
    ("개념보고서", "7.2.2.8", "자료 출처와 가정(예시값 표시 포함)", "1_입력 '근거' 열 · 2_시험자료"),
    ("정량 결과", "7.3.3.4.1", "빈도·확률 결과(기간 단위 /동·년 명시)", "3_사건수목"),
    ("정량 결과", "7.3.3.4.2", "결과 분석: 군집별 결과와 대표성(반복시험 가중, 발화층 균등)", "4_결과모델"),
    ("정량 결과", "7.3.3.4.3", "위험도 = Σ F·C, 시나리오가 전체 스펙트럼을 대표하는 이유", "5_위험도"),
    ("정량 결과", "7.3.3.4.4", "불확실성: 인식론적·우연적 구분, 전파 방법, 민감도", "6_불확실성(MC)"),
    ("정량 결과", "7.3.3.4.5", "소프트웨어·모델 평가: 버전, 검증(V&V), 적용 범위", "7_검증 · 보고서 2.5"),
    ("정량 결과", "7.3.3.4.6", "한계: 3 m 시험 → 실물 외삽(s), 통계 부족, 독립 가정 등", "보고서 2.6"),
    ("정량 결과", "7.3.3.4.7", "결론: 허용기준과 비교(판정과 신뢰수준)", "5_위험도 · 6_불확실성(MC)"),
    ("정량 결과", "7.3.3.4.8", "참고문헌: 도면·시험성적서·통계·기준의 판/연도", "보고서 2.8"),
    ("운영·유지관리", "7.4.2", "한계·가정 목록과 이를 지키는 관리 대책", "보고서 3"),
    ("운영·유지관리", "7.4.3", "변경 관리: 외장 교체·리모델링·차단재 훼손 시 재평가", "보고서 3"),
    ("운영·유지관리", "7.4.4", "점검·시험·유지관리(ITM): 스프링클러, 배리어·차단재, 필로티 가연물", "보고서 3"),
    ("검토 질문", "8.3 (1)", "평가 목적이 명시되었는가", "보고서 1.2"),
    ("검토 질문", "8.3 (2)", "평가 범위가 명시되었는가", "보고서 1.2"),
    ("검토 질문", "8.3 (3)", "방법과 그 적합성 근거가 있는가", "보고서 1.7"),
    ("검토 질문", "8.3 (4)", "한계가 명시되었는가", "보고서 2.6"),
    ("검토 질문", "8.3 (5)", "분석 결과가 포함되었는가", "5_위험도 · 6_불확실성(MC)"),
    ("검토 질문", "8.3 (6)", "결론이 포함되었는가", "보고서 2.7"),
    ("검토 질문", "8.3 (7)", "운영·유지관리 매뉴얼이 있는가", "보고서 3"),
    ("검토 질문", "8.3 (8)", "변경 관리 지침이 있는가", "보고서 3"),
    ("검토 질문", "8.3 (9)", "점검·시험·유지관리 프로그램이 있는가", "보고서 3"),
]


def checklist(wb):
    ws = wb.create_sheet("8_문서화점검")
    title(ws, "⑦ 문서화·검토 점검표 — NFPA 551 7.2(개념보고서), 7.3.3.4(정량 모델 문서), 7.4(운영·유지관리), 8.3(검토 질문)",
          "'위치'는 이 계산서의 시트와 함께 드리는 서술 양식(외벽_QFRA_서술양식.docx)의 절 번호입니다. 상태 칸은 목록에서 고릅니다.")
    header(ws, 4, ["구분", "조항", "확인 내용", "위치", "상태", "메모"], widths=[13, 11, 62, 30, 12, 40])
    dv = DataValidation(type="list", formula1='"완료,작성 중,미작성,해당 없음"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (a, b, c, d) in enumerate(CHECK, start=5):
        rc(ws, i, 1, a, font=F_BOLD)
        rc(ws, i, 2, b, align=CENTER)
        rc(ws, i, 3, c)
        rc(ws, i, 4, d)
        rc(ws, i, 5, "미작성", font=F_INPUT, fill=FILL_TODO, align=CENTER)
        dv.add(f"E{i}")
        rc(ws, i, 6, "", fill=FILL_TODO)
    last = 4 + len(CHECK)
    rc(ws, last + 2, 3, "완료 항목 수", font=F_BOLD)
    rc(ws, last + 2, 4, f'=COUNTIF(E5:E{last},"완료")&" / "&COUNTA(C5:C{last})', align=CENTER, fill=FILL_KEY)


def main():
    params = pd.read_csv(BASE / "inputs" / "parameters.csv")
    tests = pd.read_csv(BASE / "inputs" / "test_inputs.csv")
    point = pd.read_csv(BASE / "results" / "point_estimate.csv")
    mc = pd.read_csv(BASE / "results" / "mc_summary.csv")
    ex = pd.read_csv(BASE / "results" / "exceedance.csv")
    sens = pd.read_csv(BASE / "results" / "sensitivity_rank.csv")
    wb = Workbook()
    readme(wb)
    inputs(wb, params)
    test_rows, summary = tests_sheet(wb, tests)
    event_tree(wb)
    consequences(wb, tests, test_rows, summary)
    rowmap = risk_sheet(wb, point)
    mc_sheet(wb, mc, ex, sens, params)
    verify_sheet(wb, point, rowmap)
    checklist(wb)
    for ws in wb.worksheets:                      # print: landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
