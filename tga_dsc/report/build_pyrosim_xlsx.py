#!/usr/bin/env python3
"""PyroSim (FDS) material input workbook from the TGA results: report/PyroSim_입력값_정리.xlsx

Sheets: 읽는 법 (legend, sources) · MATL 입력값 (one column per material; density and specific heat come from the
measurement sheets once filled in) · 반응 파라미터 (TGA stages; REFERENCE_RATE, E and A as formulas of the inputs) ·
밀도 측정 · 비열 측정 (templates to fill in) · 측정 곡선 (mean measured curves for an FDS TGA_ANALYSIS comparison).

    python ../pyrosim_inputs.py && python build_pyrosim_xlsx.py
    python <xlsx skill>/scripts/recalc.py PyroSim_입력값_정리.xlsx      # fills the formula values
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
RES = HERE.parent / "results"
OUT = HERE / "PyroSim_입력값_정리.xlsx"

ORDER = ["EPS_C", "EPS_SC", "PIR_C", "PIR_SC", "PF_out", "PF_in"]
LABEL = {"EPS_C": "EPS 가연성", "EPS_SC": "EPS 준불연", "PIR_C": "PIR(PUR) 가연성", "PIR_SC": "PIR(PUR) 준불연",
         "PF_out": "PF 준불연 외부용", "PF_in": "PF 준불연 내부용"}
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, color="4B5563")
F_INPUT = Font(name=FONT, size=10, color="0000FF")      # hardcoded inputs / measured values
F_LINK = Font(name=FONT, size=10, color="008000")       # links to another sheet
FILL_HEAD = PatternFill("solid", fgColor="23303D")
FILL_TODO = PatternFill("solid", fgColor="FFF2CC")       # cells to fill in
FILL_EX = PatternFill("solid", fgColor="F3F4F6")         # example row
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def cell(ws, ref, value, font=F_BASE, fill=None, fmt=None, align=WRAP, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align
    if border:
        c.border = BOX
    return c


def header(ws, row, names, widths=None):
    for j, n in enumerate(names, start=1):
        cell(ws, f"{get_column_letter(j)}{row}", n, F_HEAD, FILL_HEAD, align=CENTER)
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]


def readme(wb):
    ws = wb.active
    ws.title = "읽는 법"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    cell(ws, "A1", "PyroSim(FDS) 재료 입력값 정리 — 단열재 6종", F_TITLE, border=False, align=Alignment(vertical="center"))
    rows = [
        ("목적", "TGA–DSC 결과를 PyroSim 재료(MATL) 입력에 옮길 때, 이 시험으로 구한 값과 따로 측정해야 하는 값을 나눠 정리합니다."),
        ("핵심", "밀도·비열·열전도율·반응열은 이 TGA–DSC 자료로 구할 수 없습니다(시료 초기 질량·블랭크·열류 교정 없음). "
                 "TGA는 열분해 반응의 기준온도·기준속도·단계 수에 씁니다."),
        ("칸 색", "노란 바탕 = 직접 입력할 칸 · 파란 글자 = 입력값/측정값(가정 포함) · 검은 글자 = 계산식 · 초록 글자 = 다른 시트에서 가져온 값"),
        ("시트", "MATL 입력값 → 재료별 입력 항목 / 반응 파라미터 → TGA 단계별 REFERENCE_TEMPERATURE·REFERENCE_RATE와 FDS가 계산하는 A·E / "
                 "밀도 측정·비열 측정 → 측정 기록(채우면 MATL 입력값에 평균이 들어감) / 측정 곡선 → FDS TGA_ANALYSIS 결과와 비교할 곡선"),
        ("REFERENCE_RATE", "FDS 정의: 기준온도에서 '반응하는 성분'의 정규화 질량감소속도(1/s). 이 자료의 전환율 α는 790 °C까지 감소량으로 "
                           "정규화했으므로, 감소량의 몫 f를 차지하는 단계의 값은 승온속도 × 피크 속도 ÷ f 입니다(초기 질량 불필요)."),
        ("A·E 계산식", "FDS read.f90: E = e·r_p·R·T_p²/β, A = e·r_p·exp(E/(R·T_p)) (1차 반응 가정, β = HEATING_RATE). "
                       "REFERENCE_TEMPERATURE를 쓰면 FDS가 이 식으로 A·E를 계산하므로 A·E를 따로 넣지 않습니다."),
        ("산화 반응", "탄화층 산화를 N_O2 = 1(산소 의존)로 넣으면 FDS 반응속도에 산소 몰분율이 곱해집니다. 공기 중 TGA로 구한 값이면 "
                      "A를 X_O2(공기 0.209)로 나눠 A·E를 직접 입력합니다. O 가스가 순수 산소였다면 나누지 않습니다(가스 종류 확인 필요)."),
        ("출처", "FDS User's Guide 'Estimating Kinetic Parameters' · 'Thermal Properties' · 'Simulating Bench-scale Measurements "
                 "(TGA_ANALYSIS)' (github.com/firemodels/fds, Manuals/FDS_User_Guide) · FDS Source/read.f90 · "
                 "이 저장소 tga_dsc/pyrosim_inputs.py, results/pyrosim_reactions.csv"),
        ("주의", "카탈로그 k·ρ·c를 그대로 넣은 맨 단열재는 중공층 시험보다 훨씬 빨리 착화합니다(기존 역산: 모델 16–51 s, 시험 188–288 s). "
                 "면재·마감층을 함께 모델링하고, 중공층 시험(착화 시점·t600)으로 검증하십시오."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        cell(ws, f"A{i}", k, F_BOLD)
        cell(ws, f"B{i}", v)
        ws.row_dimensions[i].height = 44


def reactions(wb, st):
    ws = wb.create_sheet("반응 파라미터")
    cell(ws, "A1", "TGA 단계별 반응 파라미터 (두 반복 평균, 5 °C/min)", F_TITLE, border=False, align=Alignment(vertical="center"))
    cell(ws, "A3", "HEATING_RATE (K/min)", F_BOLD)
    cell(ws, "B3", 5, F_INPUT, fmt="0")
    ws["B3"].comment = Comment("TGA 승온속도 = FDS HEATING_RATE. 원자료의 실측 승온속도는 4.94–4.98 K/min.", "pyrosim_inputs")
    cell(ws, "A4", "기체상수 R (J/(mol·K))", F_BOLD)
    cell(ws, "B4", 8.314462618, F_INPUT, fmt="0.000")
    cell(ws, "A5", "X_O2 (O 가스가 공기일 때)", F_BOLD)
    cell(ws, "B5", 0.209, F_INPUT, fmt="0.000")
    ws["B5"].comment = Comment("원자료에 O 가스 종류가 없음. 공기라고 가정한 값이며, 순수 산소였다면 1.0으로 바꿉니다.", "pyrosim_inputs")
    cols = ["재료", "분위기", "단계", "PyroSim 용도", "구간 시작 (°C)", "구간 끝 (°C)", "손실 몫 f", "단계 완료",
            "REFERENCE_TEMPERATURE (°C)", "피크 속도 (%/°C)", "REFERENCE_RATE (1/s)", "두 반복 REFERENCE_RATE (1/s)",
            "E (kJ/mol)", "A (1/s)", "A ÷ X_O2 (산화, N_O2 = 1)", "자료 판정"]
    widths = [17, 9, 6, 30, 10, 10, 9, 8, 15, 11, 14, 20, 10, 12, 16, 11]
    header(ws, 7, cols, widths)
    verdict_ko = {"reliable": "신뢰", "conditional": "조건부", "endpoint": "종점 의존", "remeasure": "재측정"}
    rowmap = {}
    for i, r in enumerate(st.itertuples(), start=8):
        rowmap[(r.material, r.atmosphere, r.stage)] = i
        todo = None
        cell(ws, f"A{i}", r.label, F_BOLD)
        cell(ws, f"B{i}", r.atmosphere, align=CENTER)
        cell(ws, f"C{i}", r.stage, align=CENTER)
        cell(ws, f"D{i}", r.role)
        cell(ws, f"E{i}", r.T_from_C, fmt="0", align=CENTER)
        cell(ws, f"F{i}", r.T_to_C, fmt="0", align=CENTER)
        cell(ws, f"G{i}", r.loss_share, F_INPUT, fmt="0.000", align=CENTER)
        cell(ws, f"H{i}", "예" if r.finished else "아니오", align=CENTER)
        cell(ws, f"I{i}", r.REFERENCE_TEMPERATURE_C, F_INPUT, todo, fmt="0", align=CENTER)
        cell(ws, f"J{i}", r.peak_rate_pct_per_C, F_INPUT, fmt="0.000", align=CENTER)
        cell(ws, f"K{i}", f"=($B$3/60)*(J{i}/100)/G{i}", fmt="0.00E+00", align=CENTER)
        cell(ws, f"L{i}", r.REFERENCE_RATE_runs, align=CENTER)
        cell(ws, f"M{i}", f"=EXP(1)*K{i}*$B$4*(I{i}+273.15)^2/($B$3/60)/1000", fmt="0", align=CENTER)
        cell(ws, f"N{i}", f"=EXP(1)*K{i}*EXP(M{i}*1000/($B$4*(I{i}+273.15)))", fmt="0.00E+00", align=CENTER)
        cell(ws, f"O{i}", f"=N{i}/$B$5" if r.atmosphere == "공기(O)" else "—", fmt="0.00E+00", align=CENTER)
        cell(ws, f"P{i}", verdict_ko[r.tga_verdict], align=CENTER)
        if not r.finished:
            for col in "IJK":
                ws[f"{col}{i}"].font = Font(name=FONT, size=10, color="9CA3AF")
            ws[f"I{i}"].comment = Comment("790 °C에서도 이 단계의 감소가 진행 중이라 피크가 범위 끝에 걸림. 기준값으로 쓰지 말 것.", "pyrosim_inputs")
    last = 7 + len(st)
    notes = [
        "파란 글자(손실 몫 f, 기준온도, 피크 속도)는 pyrosim_inputs.py가 원자료에서 구한 값입니다. REFERENCE_RATE·E·A는 이 값들의 수식입니다.",
        "회색 글자 = 790 °C에서 끝나지 않은 단계(피크가 측정 범위 끝). 800 °C 등온 유지 또는 1000 °C까지 측정 후 다시 구해야 합니다.",
        "PF 공기 2단계는 열분해와 탄화층 산화가 겹친 구간(f ≈ 0.9)이라 산화만의 속도는 이보다 큽니다. FDS TGA_ANALYSIS로 공기 곡선과 맞춰 조정하십시오.",
        "EPS 가연성 공기는 재측정 대상(질량감소 0.16·0.30 mg)이므로 값을 쓰지 마십시오.",
    ]
    for k, t in enumerate(notes):
        cell(ws, f"A{last + 2 + k}", t, F_NOTE, border=False, align=Alignment(vertical="center"))
    ws.freeze_panes = "B8"
    return rowmap


def matl(wb, rowmap):
    ws = wb.create_sheet("MATL 입력값", 1)
    cell(ws, "A1", "재료별 PyroSim(FDS) MATL 입력값 — 노란 칸은 실측 또는 제품 시험성적서 값으로 채움", F_TITLE, border=False,
         align=Alignment(vertical="center"))
    cols = ["PyroSim 항목", "FDS 키워드", "단위", "구하는 방법"] + [LABEL[m] for m in ORDER] + ["비고"]
    header(ws, 3, cols, [22, 24, 11, 40, 14, 14, 14, 14, 14, 14, 44])
    mcol = {m: get_column_letter(5 + j) for j, m in enumerate(ORDER)}
    rp = "'반응 파라미터'"
    lines = []

    def add(item, kw, unit, how, values, note="", kind="todo", fmt=None):
        lines.append((item, kw, unit, how, values, note, kind, fmt))

    add("밀도", "DENSITY", "kg/m³", "보드 실측 (KS M ISO 845) — '밀도 측정' 시트의 재료별 평균",
        {m: f"=IFERROR(AVERAGEIFS('밀도 측정'!$G$6:$G$80,'밀도 측정'!$A$6:$A$80,{mcol[m]}$3),\"\")" for m in ORDER},
        "측정 전 임시 범위(기존 분석의 카탈로그 범위): 15–60 kg/m³ — 제품 시험성적서 값 우선", "formula", "0.0")
    add("비열", "SPECIFIC_HEAT (또는 SPECIFIC_HEAT_RAMP)", "kJ/(kg·K)", "사파이어법 DSC (ASTM E1269 / ISO 11357-4) — '비열 측정' 시트의 재료별 평균",
        {m: f"=IFERROR(AVERAGEIFS('비열 측정'!$I$6:$I$80,'비열 측정'!$A$6:$A$80,{mcol[m]}$3),\"\")" for m in ORDER},
        "임시 범위: 1.3–1.5 kJ/(kg·K). 온도에 따라 변하므로 측정값은 SPECIFIC_HEAT_RAMP(온도별)로 넣는 것이 좋음", "formula", "0.00")
    add("열전도율", "CONDUCTIVITY", "W/(m·K)", "제품 시험성적서 (KS L 9016 / ISO 8301)", {m: None for m in ORDER},
        "임시 범위: 0.018–0.040 W/(m·K). 고온 값은 CONDUCTIVITY_RAMP", "todo", "0.000")
    add("방사율", "EMISSIVITY", "–", "FDS 기본값", {m: 0.9 for m in ORDER}, "측정값이 없으면 기본값 0.9", "input", "0.0")
    add("반응 수", "N_REACTIONS", "–", "TGA DTG 단계 (반응 1 열분해, 반응 2 탄화층 산화)",
        {"EPS_C": 1, "EPS_SC": 1, "PIR_C": 2, "PIR_SC": 2, "PF_out": 2, "PF_in": 2}, "EPS는 한 단계, PIR(PUR)·PF는 열분해 + 탄화층 산화",
        "input", "0")
    add("승온속도", "HEATING_RATE(:)", "K/min", "TGA 조건", {m: f"={rp}!$B$3" for m in ORDER}, "모든 반응에 같은 값", "link", "0")
    add("반응 1 기준온도", "REFERENCE_TEMPERATURE(1)", "°C", "TGA N₂ DTG 피크 (반응 1)",
        {m: f"={rp}!I{rowmap[(m, 'N₂', 1)]}" for m in ORDER}, "PIR(PUR)·PF는 1단계 피크. PF는 피크가 넓어 정확도가 낮음", "link", "0")
    add("반응 1 기준속도", "REFERENCE_RATE(1)", "1/s", "승온속도 × 피크 속도 ÷ 단계 몫 (N₂)",
        {m: f"={rp}!K{rowmap[(m, 'N₂', 1)]}" for m in ORDER}, "REFERENCE_TEMPERATURE와 함께 쓰고 A·E는 비워 둠", "link", "0.00E+00")
    add("반응 1 잔류물 수율", "NU_MATL(1,1)", "kg/kg", "790 °C 잔류 질량 ÷ 초기 질량 — 시료 초기 질량 기록 필요",
        {m: None for m in ORDER}, "현재 자료로 계산 불가(질량 영점 없음). EPS 가연성은 잔류물이 거의 없을 것으로 보이나 확인 필요", "todo", "0.00")
    add("반응 1 반응열", "HEAT_OF_REACTION(1)", "kJ/kg", "교정된 DSC(블랭크 + 열류 교정 + 초기 질량) 또는 문헌",
        {m: None for m in ORDER}, "부호만 확인됨: EPS 가연성 N₂ 398 °C 흡열(+). 양수 = 흡열", "todo", "0")
    add("반응 2 기준온도", "REFERENCE_TEMPERATURE(2)", "°C", "TGA 공기 탄화층 산화 피크 (반응 2)",
        {m: (f"={rp}!I{rowmap[(m, '공기(O)', 2)]}" if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "N_O2 = 1로 넣을 때는 아래 A·E를 직접 입력", "link", "0")
    add("반응 2 기준속도", "REFERENCE_RATE(2)", "1/s", "승온속도 × 피크 속도 ÷ 단계 몫 (공기)",
        {m: (f"={rp}!K{rowmap[(m, '공기(O)', 2)]}" if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "PF 공기 값은 열분해와 겹친 구간 기준 — 산화만의 속도는 더 큼", "link", "0.00E+00")
    add("반응 2 산소 차수", "N_O2(2)", "–", "탄화층 산화 = 산소 의존", {m: (1 if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "", "input", "0")
    add("반응 2 A (N_O2 = 1)", "A(2)", "1/s", "'반응 파라미터'의 A ÷ X_O2",
        {m: (f"={rp}!O{rowmap[(m, '공기(O)', 2)]}" if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "O 가스가 공기일 때. 순수 산소였다면 '반응 파라미터'의 A 그대로", "link", "0.00E+00")
    add("반응 2 E", "E(2)", "J/mol", "'반응 파라미터'의 E × 1000",
        {m: (f"={rp}!M{rowmap[(m, '공기(O)', 2)]}*1000" if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "", "link", "0")
    add("반응 2 반응열", "HEAT_OF_REACTION(2)", "kJ/kg", "교정 DSC 또는 문헌 (산화 = 발열, 음수)", {m: (None if m not in ("EPS_C", "EPS_SC") else "—") for m in ORDER},
        "부호만 확인됨: PF 내부용 공기 474 °C 발열", "todo", "0")
    add("연소열", "HEAT_OF_COMBUSTION", "kJ/kg", "MCC (ASTM D7309) 또는 콘칼로리미터 (ISO 5660-1)", {m: None for m in ORDER},
        "DSC로는 구할 수 없음", "todo", "0")
    for i, (item, kw, unit, how, values, note, kind, fmt) in enumerate(lines, start=4):
        cell(ws, f"A{i}", item, F_BOLD)
        cell(ws, f"B{i}", kw)
        cell(ws, f"C{i}", unit, align=CENTER)
        cell(ws, f"D{i}", how)
        for m in ORDER:
            v = values[m]
            font = {"link": F_LINK, "input": F_INPUT}.get(kind, F_BASE)
            fill = FILL_TODO if (kind in ("todo", "formula") and v != "—") else None
            cell(ws, f"{mcol[m]}{i}", v, font, fill, fmt=fmt, align=CENTER)
        cell(ws, f"K{i}", note, F_NOTE)
        ws.row_dimensions[i].height = 34
    n = 4 + len(lines)
    cell(ws, f"A{n + 1}", "노란 칸 = 채울 칸(밀도·비열은 측정 시트를 채우면 자동으로 평균이 들어감). 초록 = '반응 파라미터'에서 가져온 값. "
                         "REFERENCE_TEMPERATURE를 쓰는 반응에는 A·E를 넣지 않습니다(N_O2 = 1인 반응 2만 A·E 직접 입력).",
         F_NOTE, border=False, align=Alignment(vertical="center"))
    ws.freeze_panes = "E4"


def measurements(wb):
    labels = [LABEL[m] for m in ORDER]
    dv_list = '"' + ",".join(labels + ["예시"]) + '"'
    # density
    ws = wb.create_sheet("밀도 측정")
    cell(ws, "A1", "밀도 측정 기록 (KS M ISO 845: 겉보기 밀도, 시편 3개 이상, 23 °C · 50 %RH 상태 조절)", F_TITLE, border=False,
         align=Alignment(vertical="center"))
    cell(ws, "A2", "노란 칸에 입력 → G열 밀도가 계산되고 'MATL 입력값'에 재료별 평균이 들어감. 5행은 형식을 보여 주는 예시(재료 '예시'는 평균에서 제외).",
         F_NOTE, border=False, align=Alignment(vertical="center"))
    header(ws, 4, ["재료", "시편", "가로 (mm)", "세로 (mm)", "두께 (mm)", "질량 (g)", "밀도 (kg/m³)", "비고"], [18, 7, 11, 11, 11, 11, 13, 40])
    ex = ["예시", 1, 100.0, 100.0, 50.0, 9.0]
    for j, v in enumerate(ex):
        cell(ws, f"{get_column_letter(j + 1)}5", v, F_INPUT, FILL_EX, align=CENTER)
    cell(ws, "G5", "=IF(COUNT(C5:F5)=4,F5/1000/(C5*D5*E5/1000000000),\"\")", fill=FILL_EX, fmt="0.0", align=CENTER)
    cell(ws, "H5", "면재가 붙은 제품은 면재를 떼어낸 심재와 전체를 따로 기록", F_NOTE, FILL_EX)
    for i in range(6, 81):
        for col in "ABCDEF":
            cell(ws, f"{col}{i}", None, F_INPUT, FILL_TODO, align=CENTER)
        cell(ws, f"G{i}", f"=IF(COUNT(C{i}:F{i})=4,F{i}/1000/(C{i}*D{i}*E{i}/1000000000),\"\")", fmt="0.0", align=CENTER)
        cell(ws, f"H{i}", None)
    dv = DataValidation(type="list", formula1=dv_list, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A6:A80")
    ws.freeze_panes = "A5"
    # specific heat
    ws = wb.create_sheet("비열 측정")
    cell(ws, "A1", "비열 측정 기록 — 사파이어법 DSC (ASTM E1269 / ISO 11357-4): 빈 도가니 · 사파이어 · 시료를 같은 조건으로 측정",
         F_TITLE, border=False, align=Alignment(vertical="center"))
    cell(ws, "A2", "c_p,시료 = c_p,사파이어 × (m_사파이어 / m_시료) × (HF_시료 − HF_블랭크) / (HF_사파이어 − HF_블랭크). "
                   "사파이어 c_p는 표준(ASTM E1269 표 또는 NIST SRM 720 인증서)의 값을 온도별로 입력. 5행은 형식 예시.",
         F_NOTE, border=False, align=Alignment(vertical="center"))
    header(ws, 4, ["재료", "온도 (°C)", "블랭크 열류 (mW)", "사파이어 열류 (mW)", "시료 열류 (mW)", "사파이어 질량 (mg)",
                   "시료 질량 (mg)", "사파이어 c_p (J/(g·K))", "시료 c_p (kJ/(kg·K))", "비고"], [18, 10, 13, 14, 12, 13, 12, 15, 15, 36])
    ex = ["예시", 100, 1.00, 4.80, 2.70, 25.0, 6.0, 0.85]
    for j, v in enumerate(ex):
        cell(ws, f"{get_column_letter(j + 1)}5", v, F_INPUT, FILL_EX, align=CENTER)
    formula = lambda i: (f"=IF(AND(COUNT(B{i}:H{i})=7,D{i}<>C{i},G{i}<>0),H{i}*(F{i}/G{i})*(E{i}-C{i})/(D{i}-C{i}),\"\")")  # noqa: E731
    cell(ws, "I5", formula(5), fill=FILL_EX, fmt="0.00", align=CENTER)
    cell(ws, "J5", "예시 값(형식만). 측정 범위는 분해 시작 전(약 30–150 °C)", F_NOTE, FILL_EX)
    for i in range(6, 81):
        for col in "ABCDEFGH":
            cell(ws, f"{col}{i}", None, F_INPUT, FILL_TODO, align=CENTER)
        cell(ws, f"I{i}", formula(i), fmt="0.00", align=CENTER)
        cell(ws, f"J{i}", None)
    dv = DataValidation(type="list", formula1=dv_list, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A6:A80")
    ws.freeze_panes = "A5"


def curves(wb):
    ws = wb.create_sheet("측정 곡선")
    with open(RES / "pyrosim_tga_curves.csv", encoding="utf-8") as fh:
        fh.readline()
        cv = pd.read_csv(fh)
    cell(ws, "A1", "측정 곡선 (두 반복 평균) — FDS TGA_ANALYSIS 결과(CHID_tga.csv)와 비교용", F_TITLE, border=False,
         align=Alignment(vertical="center"))
    cell(ws, "A2", "질량(정규화) = 1 − α, α = 승온 시작 이후 감소량 ÷ 790 °C까지 감소량. 속도 = dα/dt (1/s, 5 K/min). "
                   "FDS 결과의 질량 Y(초기 질량 기준)는 α_FDS = (1 − Y) / (1 − Y(790 °C))로 바꿔 비교.", F_NOTE, border=False,
         align=Alignment(vertical="center"))
    names = ["온도 (°C)"]
    keys = ["T_C"]
    for m in ORDER:
        for atm, ko in (("N", "N₂"), ("O", "공기")):
            names += [f"{LABEL[m]} {ko} 질량", f"{LABEL[m]} {ko} 속도 (1/s)"]
            keys += [f"{m}_{atm}_mass_norm", f"{m}_{atm}_rate_per_s"]
    header(ws, 4, names, [9] + [13] * (len(names) - 1))
    ws.row_dimensions[4].height = 44
    for i, row in enumerate(cv[keys].itertuples(index=False), start=5):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=float(v))
            c.font = F_BASE
            c.number_format = "0" if j == 1 else ("0.0000" if (j % 2 == 0) else "0.00E+00")
    ws.freeze_panes = "B5"


def main():
    st = pd.read_csv(RES / "pyrosim_reactions.csv")
    wb = Workbook()
    readme(wb)
    rowmap = reactions(wb, st)
    matl(wb, rowmap)
    measurements(wb)
    curves(wb)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
