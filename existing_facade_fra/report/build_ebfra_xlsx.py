#!/usr/bin/env python3
"""기축 건축물 외벽 화재위험성 평가 계산서: report/기축건축물_외벽화재위험성_평가계산서.xlsx

The building sheets compute the method of ebfra.py with cell formulas: zone table → material values by evidence level
(type default, TGA prediction, test value) → per-zone helper table by spread distance j → origin storey i × j grids
(IE1 with and without the façade path) and rows for fires at the foot of the zone (IE2) and from the zone below (IE3)
→ R, R_ref, IR, r, grade. Remediation switches (Y/N) on each sheet reproduce the options of ebfra.py. Monte Carlo,
evidence-level spread and sensitivity come from ebfra.py as values.

    python ../ebfra.py && python build_ebfra_xlsx.py
    python <xlsx skill>/scripts/recalc.py 기축건축물_외벽화재위험성_평가계산서.xlsx
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = HERE / "기축건축물_외벽화재위험성_평가계산서.xlsx"
D = json.loads((HERE / "ebfra_data.json").read_text(encoding="utf-8"))

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_SMALL = Font(name=FONT, size=8)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, color="4B5563")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_KEY = Font(name=FONT, size=11, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="23303D")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2F3E2")
FILL_GREY = PatternFill("solid", fgColor="F3F4F6")
FILL_WARN = PatternFill("solid", fgColor="FDE2E2")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SCI = "0.00E+00"
NZ = 4            # zones per building sheet
NS = 30           # storeys
NJ = NS - 1       # spread distance j = 1 … 29
Z_IN, Z_CALC, K0 = 19, 26, 33          # header rows: zone inputs, zone calculation, storey table
G0, BLOCK = 66, 105                    # first grid block row, rows per zone block


def cell(ws, ref, value, font=F_BASE, fill=None, fmt=None, align=WRAP, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align
    if border:
        c.border = BOX
    return c


def rc(ws, row, col, value, **kw):
    return cell(ws, f"{get_column_letter(col)}{row}", value, **kw)


def header(ws, row, names, col0=1, widths=None, height=None):
    for k, n in enumerate(names):
        rc(ws, row, col0 + k, n, font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        if widths:
            ws.column_dimensions[get_column_letter(col0 + k)].width = widths[k]
    if height:
        ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    cell(ws, "A1", text, F_TITLE, border=False, align=Alignment(vertical="center"))
    if sub:
        cell(ws, "A2", sub, F_NOTE, border=False, align=Alignment(vertical="center"))


def band(ws, row, text, ncol, col0=1):
    rc(ws, row, col0, text, font=F_SUB, fill=FILL_GREY)
    ws.merge_cells(start_row=row, start_column=col0, end_row=row, end_column=col0 + ncol - 1)


def note(ws, row, text, col=1):
    rc(ws, row, col, text, font=F_NOTE, border=False, align=Alignment(vertical="center"))


def L(c):
    return get_column_letter(c)


def define(wb, name, sheet, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{ref}")


# ------------------------------------------------------------------------------------------------ 읽는 법
def readme(wb):
    ws = wb.active
    ws.title = "읽는 법"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    title(ws, "기축 건축물 외벽 화재위험성 평가 계산서 — 건물 형태와 관계없이 '외벽 구간'으로 나누어 계산",
          "연구 방법론 제안용 · 예제 건물 3동(가상) · 빈도·분기확률 등은 예시값(4_공통변수의 근거 열 참조)")
    rows = [
        ("무엇을 계산하나", "외벽을 통해 번지는 화재로 1년에 노출되는 재실자 수 R(명/년). 재실자 1인당 IR = R / 재실자, 자택 화재 위험(λ_res)에 대한 비 r = IR / λ_res로 5등급을 매깁니다."),
        ("평가식", "0_평가식 시트에 식·기호·계산 위치를 모았습니다. R = Σ구간 [ Σ발화층 f_i Σj c·(외벽 경로·창→창 도약) + 외벽 하단 화원 + 하부 구간 분출화염 ]"),
        ("건물 형태 다루기", "입면을 '같은 외벽 시스템·같은 노출'인 구간으로 나눕니다(층 범위 b–t, 수직 열 1·2). 필로티, 저층 상가 + 고층 주거, 측벽, 커튼월 등 어떤 형태든 구간의 조합으로 입력합니다."),
        ("재료 인자", "외벽 유형별 t₆₀₀(층당 확산시간)·p_B(자립 연소)·A(급발). 증거 수준 E0(사진·허가연도) → E1(현장) → E2(시료: FT-IR·콘·TGA) → E3(시스템 시험·시뮬레이션)으로 바뀝니다(2_외벽유형)."),
        ("색", "노란 칸 = 입력(파란 글자), 흰 칸 = 수식, 연두 칸 = 핵심 결과, 회색 = 구분선·참고값."),
        ("내 건물 평가", "8_빈양식에 건물 정보(B6–B13)와 구간(행 20–23)을 넣으면 결과(I5–I16)가 바로 나옵니다. 대책 스위치(F5–F9)를 Y로 바꾸면 대책 후 위험도가 계산됩니다."),
        ("예제", "5_예제A 필로티 다세대(드라이비트), 6_예제B 15층 아파트(중공층 PIR 준불연), 7_예제C 25층 주상복합(복합패널) — 모두 가상 건물입니다."),
        ("Python 값", "9_대책비교·10_증거수준·11_불확실성은 ebfra.py(몬테카를로 4000회 등)의 결과 값입니다. 12_검증에서 엑셀 수식과 파이썬 값이 같은지 확인합니다."),
        ("재현", "python existing_facade_fra/ebfra.py → python existing_facade_fra/report/build_ebfra_xlsx.py → recalc(LibreOffice)."),
        ("주의", "빈도·분기확률·시간은 절차를 보이기 위한 예시값입니다. 허가 시점 기준(3_허가시점기준)의 연도 일부는 법령 원문으로 확정해야 합니다."),
    ]
    for k, (a, b) in enumerate(rows, start=4):
        rc(ws, k, 1, a, font=F_BOLD, fill=FILL_GREY)
        rc(ws, k, 2, b)
        ws.row_dimensions[k].height = 32


# ------------------------------------------------------------------------------------------------ 0 평가식
def formula_sheet(wb):
    ws = wb.create_sheet("0_평가식")
    title(ws, "평가식 — 식, 기호, 이 계산서에서 계산하는 곳", "구간 z · 발화층 i · 확산 층수 j · 대상 층 k = i + j")
    header(ws, 4, ["항목", "식", "의미", "계산 위치"], widths=[20, 70, 52, 30])
    rows = [
        ("건물 위험도", "R = Σ_z ( R1_z + R2_z + R3_z )", "외벽을 통해 1년에 노출되는 재실자 수(명/년)", "건물 시트 I5"),
        ("① 실내 화재 → 창 분출", "R1_z = Σ_i f_i Σ_j c(k) [ p_e·a·(b·q_min + (1−b)·q_a) + (1 − p_e·a)·b·q_b ]",
         "구간 뒤 세대·실의 화재가 창으로 분출 → 외벽 경로 a 또는 창→창 도약 b로 위층에 닿음", "구간 블록 'IE1 격자', 계산표 Y열"),
        ("② 외벽 하단 화원", "R2_z = Σ_s F_s Σ_j c(k) p_e2 a_s(j) q_s(j)", "필로티 차량·쓰레기·실외기 등 외벽 바로 아래 화재", "구간 블록 'IE2' 행, 계산표 AA열"),
        ("③ 하부 구간 분출화염", "R3_z = f_하부 Σ_j c(k) p_e2 a(j) [ b q_min + (1−b) q_a − b q_b ]", "아래 구간 최상층 창의 화염이 이 구간 하단을 가열(예: 저층 상가 → 고층 패널)", "구간 블록 'IE3' 행, 계산표 AB열"),
        ("발화 빈도", "f_i = (n_u·λ_res 또는 면적·λ_com) · (1 − p_s·스프링클러_i) · p_FO · p_v · 창계수", "층마다 외벽에 분출화염이 생기는 빈도(/년)", "계산표 N열 × 층별 스프링클러"),
        ("외벽 경로", "a(j) = (1 − p_cb)^j · S_FS(j·τ) · [ p_B + (1 − p_B)·S_e(j·τ) ]", "외벽 화재가 j개 층 위까지 오를 확률: 진압 전·노출 지속 중 또는 자립 연소", "구간 보조표 E–K·R열"),
        ("층당 확산시간", "τ = s · t₆₀₀ · h / 3 m", "3 m 중공층 시험의 t₆₀₀을 건물로 옮김(s: 실물/시험 비)", "계산표 I열"),
        ("창→창 도약", "b(j) = p_L^j · (1 − p_s)^(위층 중 스프링클러 층 수) · S_FS(j·τ_L)", "외벽이 불연이어도 남는 경로(기준 위험도 R_ref)", "구간 보조표 G·H·S열"),
        ("피난 지연", "q(t) = P(RSET > t)", "화재가 닿을 때 아직 피난하지 못한 확률(필로티 화재 + 1층 방화구획 없음: + p_st)", "구간 보조표 L–Q열"),
        ("노출 재실자", "c(k) = 확산 띠 세대 × 세대당 인원 × 실내 유입", "외벽 화재가 한 층에 닿을 때 노출되는 사람", "층별 표 E·G열"),
        ("기준 위험도", "R_ref = Σ_z Σ_i f_i Σ_j c(k)·b·q_b", "같은 건물의 외벽을 모두 불연으로 바꿨을 때(도약만)", "건물 시트 I6"),
        ("개인 위험", "IR = R / 재실자,  r = IR / λ_res", "1인당 연간 노출 확률과 자택 화재 위험(λ_res)에 대한 비", "I8 · I9"),
        ("등급", "1: r < 0.01 · 2: < 0.1 · 3: < 1 · 4: < 10 · 5: ≥ 10", "자택 화재 위험의 1 %·10 %·같은 수준·10배를 경계로", "I10"),
        ("표시", "M: C̄ = R / F_ign ≥ 10명  ·  A: 급발(용융·가연 전면 화염)", "대형 피해 가능성, 급발 연소 가능성", "I12 · I13"),
        ("보조 지표", "RR = R / R_ref,  AF = 1 − R_ref / R,  ΔR = R − R_ref", "외벽 재료 때문에 늘어난 몫(교체로 줄일 수 있는 위험)", "I14–I16"),
        ("재료 인자(E2)", "t̂₆₀₀ = 247.5 + 422.97 / r_max,air,  p̂_B = (f₄₀₀,air − 0.60)/0.10 (0.05–0.95)", "콘으로 준불연이 확인된 탄화형(f₄₀₀ ≥ 0.5) 시료만", "2_외벽유형 E2 계산기, 계산표 D·E열"),
    ]
    for k, row in enumerate(rows, start=5):
        for c, v in enumerate(row, start=1):
            rc(ws, k, c, v, font=F_BOLD if c == 1 else F_BASE)
        ws.row_dimensions[k].height = 34


# ------------------------------------------------------------------------------------------------ 1 절차
def process(wb):
    ws = wb.create_sheet("1_절차")
    title(ws, "평가 절차 — 증거가 쌓일수록 정밀해지는 5단계", "단계마다 같은 평가식으로 R·등급을 다시 계산합니다(재료 인자의 근거만 바뀜)")
    header(ws, 4, ["단계", "하는 일", "입력·자료원", "재료 인자(증거 수준)", "산출"], widths=[16, 48, 46, 34, 30])
    rows = [
        ("0 선별", "건축물대장·외관 사진으로 많은 건물을 한꺼번에 평가해 우선순위를 냄",
         "건축물대장(용도·층수·높이·허가일·구조), 로드뷰 사진, 3_허가시점기준", "E0: 사진의 외벽 유형 + 허가 시점 기준 → 보수적 기본값", "1차 등급·우선순위"),
        ("1 현장 조사", "입면을 구간으로 나누고 외벽 구성·개구부·필로티·배리어·외부 화원·소방 접근을 확인",
         "현장 조사표(외벽 구성, 창 비율, 층간 띠, 필로티 천장, 주차·쓰레기 위치, 스프링클러, 계단)", "E1: 유형 확인(등급은 표시·서류가 있을 때만)", "구간 입력 확정, 등급"),
        ("2 시료 시험", "구간 유형마다 시료를 채취해 재료를 식별하고 등급·지표를 측정",
         "FT-IR(고분자·난연제) → 콘(KS F ISO 5660-1, 준불연 판정) → TGA(공기·N₂)", "E2: 콘 가연 → 가연 경계값 / 준불연 탄화형 → TGA 예측식 / 용융형 → 급발 A", "재료 인자 확정, 등급"),
        ("3 정밀 평가", "등급이 높거나 경계에 걸린 건물: 시스템 시험·시뮬레이션",
         "중규모(3 m 중공층)·실대형(KS F 8414) 시험, FDS 건물 외벽 해석(s 보정)", "E3: 시험 t₆₀₀·자립 연소, s", "확정 등급, 대책 설계"),
        ("4 판정·대책", "등급·표시(M·A)와 대책별 위험 감소를 같은 식으로 계산해 우선순위를 정함",
         "대책 스위치(차단재·확산방지구조·하단 보호·스프링클러·교체)", "—", "등급, ΔR(회피 가능 위험), 문서(NFPA 551 7장)"),
    ]
    for k, row in enumerate(rows, start=5):
        for c, v in enumerate(row, start=1):
            rc(ws, k, c, v, font=F_BOLD if c == 1 else F_BASE)
        ws.row_dimensions[k].height = 46
    header(ws, 11, ["증거 수준", "재료 인자 출처", "t₆₀₀ 불확실성(ln 표준편차)", "근거", ""])
    ev = [("E0 · E1", "유형·허가 시점 → 보수적 기본값(준불연 미상 512 s / 가연 30 s)", 0.506, "준불연 4종의 제품 간 편차(multiscale_fra)"),
          ("E2", "TGA 예측식(준불연 탄화형) · 콘 가연이면 가연 경계값", 0.166, "TGA 예측 편차(제품 제외 교차검증)"),
          ("E3", "시스템 시험값(2회)", 0.066, "반복시험 편차 0.093 / √2")]
    for k, row in enumerate(ev, start=12):
        for c, v in enumerate(row, start=1):
            rc(ws, k, c, v, font=F_BOLD if c == 1 else F_BASE, fmt="0.000" if c == 3 else None)


# ------------------------------------------------------------------------------------------------ 2 외벽유형
def types_sheet(wb):
    ws = wb.create_sheet("2_외벽유형")
    title(ws, "외벽 유형별 재료 인자 기본값과 시료 시험(E2) 계산기",
          "시험한 시스템(E3)은 시험값, 시험 자료가 없는 시스템은 보수적 경계값(E0) — 시험이 들어오면 그 값으로 바꿉니다")
    header(ws, 4, ["키", "외벽 유형", "구성", "등급", "외벽 경로", "t₆₀₀ (s)", "p_B", "급발 A", "확산방지구조 인정", "차단재 인정", "증거", "근거"],
           widths=[14, 34, 34, 8, 9, 9, 8, 8, 10, 10, 7, 60], height=30)
    for k, t in enumerate(D["types"], start=5):
        vals = [t["key"], t["label"], t["system"], t["grade"], int(t["path"]),
                t["t600_s"] if t["t600_s"] == t["t600_s"] and t["t600_s"] is not None else None,
                t["pB"] if t["pB"] == t["pB"] and t["pB"] is not None else None, int(t["A"]),
                int(t["cb_ok"]) if t["cb_ok"] == t["cb_ok"] and t["cb_ok"] is not None else 0,
                int(t["pe_ok"]) if t["pe_ok"] == t["pe_ok"] and t["pe_ok"] is not None else 0, t["evidence"], t["basis"]]
        for c, v in enumerate(vals, start=1):
            rc(ws, k, c, v, font=F_INPUT if c in (6, 7) else F_BASE, fill=FILL_IN if c in (6, 7) and v is not None else None)
        ws.row_dimensions[k].height = 28
    last = 4 + len(D["types"])
    define(wb, "T_types", "2_외벽유형", f"$A$5:$L${last}")
    define(wb, "T_keys", "2_외벽유형", f"$A$5:$A${last}")
    note(ws, last + 1, "확산방지구조·차단재 인정 = 0: 가연재가 바깥 면이거나 패널 심재로 이어져 층간 띠·개구부 차단재로 끊기지 않음(PE 심재 복합패널, 샌드위치패널, 미상 가연재)")

    r0 = last + 3
    band(ws, r0, "E2 시료 시험 계산기 — 콘으로 등급을 가리고, 준불연 탄화형이면 TGA로 t̂₆₀₀ · p̂_B", 6)
    rows = [("콘 판정 (sc 준불연 / c 가연)", "sc", None), ("r_max,air (%/°C)", 0.5752, None), ("f₄₀₀,air (–)", 0.6288, None)]
    for k, (a, v, _) in enumerate(rows, start=r0 + 1):
        rc(ws, k, 1, a, font=F_BOLD)
        rc(ws, k, 2, v, font=F_INPUT, fill=FILL_IN)
    a, rmx, f4 = f"$B${r0 + 1}", f"$B${r0 + 2}", f"$B${r0 + 3}"
    outs = [("재료형", f'=IF({a}="c","가연(게이트)",IF({f4}<P_f_melt,"용융형","탄화형"))', None),
            ("t̂₆₀₀ (s)", f'=IF({a}="c",VLOOKUP("C_UNKNOWN",T_types,6,FALSE),IF({f4}<P_f_melt,"유형 기본값 사용",P_tga_c0+P_tga_c1/{rmx}))', "0"),
            ("p̂_B", f'=IF({a}="c",VLOOKUP("C_UNKNOWN",T_types,7,FALSE),IF({f4}<P_f_melt,"유형 기본값 사용",MIN(0.95,MAX(0.05,({f4}-P_f_lo)/(P_f_hi-P_f_lo)))))', "0.000"),
            ("급발 A", f'=IF(OR({a}="c",{f4}<P_f_melt),1,0)', "0")]
    for k, (lab, f, fmt) in enumerate(outs, start=r0 + 4):
        rc(ws, k, 1, lab, font=F_BOLD)
        rc(ws, k, 2, f, fill=FILL_KEY, fmt=fmt)
    note(ws, r0 + 8, "예시 값 = PIR(PUR) 준불연 시료(공기 5 °C/min) → t̂₆₀₀ 983 s · p̂_B 0.29. 건물 시트에서는 구간 입력 S–U열에 같은 값을 넣습니다.")


# ------------------------------------------------------------------------------------------------ 3 허가시점기준
def history_sheet(wb):
    ws = wb.create_sheet("3_허가시점기준")
    title(ws, "허가 시점의 외벽 마감재료 기준 → E0 기본 유형(0단계 선별)",
          "기준은 건축허가·신고 신청일에 적용 · '확정 필요' 표시 연도는 법령 원문(국가법령정보센터)으로 확인 후 사용")
    header(ws, 4, ["구간", "시작", "끝", "기준 대상", "요구", "E0 추정 규칙", "확인 상태", "출처"], widths=[6, 11, 11, 34, 44, 40, 22, 52], height=30)
    for k, h in enumerate(D["history"], start=5):
        vals = [h["period"], h["from"], h["to"], h["scope"], h["required"], h["e0_rule"], h["status"], h["source"]]
        for c, v in enumerate(vals, start=1):
            rc(ws, k, c, v, fill=FILL_WARN if (c == 7 and "필요" in v) else None)
        ws.row_dimensions[k].height = 44
    r0 = 11
    band(ws, r0, "E0 기본 유형 계산기", 4)
    for k, (lab, v) in enumerate([("허가(신청) 연도", 2013), ("2019년이면 11월 7일 이후 신청? (Y/N)", "N"), ("지상 층수", 5), ("높이 (m)", 14)], start=r0 + 1):
        rc(ws, k, 1, lab, font=F_BOLD)
        rc(ws, k, 2, v, font=F_INPUT, fill=FILL_IN)
        ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=2)
    y, late, n, hh = f"$B${r0 + 1}", f"$B${r0 + 2}", f"$B${r0 + 3}", f"$B${r0 + 4}"
    per = f'=IF({y}<2010,"P0",IF({y}<2015,"P1",IF(OR({y}<2019,AND({y}=2019,{late}<>"Y")),"P2","P3")))'
    rc(ws, r0 + 5, 1, "적용 구간", font=F_BOLD)
    rc(ws, r0 + 5, 2, per, fill=FILL_KEY)
    p = f"$B${r0 + 5}"
    rule = (f'=IF({p}="P0","C_UNKNOWN",IF({p}="P1",IF({n}>=30,"SC_UNKNOWN","C_UNKNOWN"),'
            f'IF({p}="P2",IF(OR({n}>=6,{hh}>=22),"SC_UNKNOWN","C_UNKNOWN"),IF(OR({n}>=3,{hh}>=9),"SC_UNKNOWN","C_UNKNOWN"))))')
    rc(ws, r0 + 6, 1, "E0 기본 유형", font=F_BOLD)
    rc(ws, r0 + 6, 2, rule, fill=FILL_KEY)
    note(ws, r0 + 7, "P3에서 5층·22 m 이하 건물은 화재확산방지구조를 두면 가연 단열재가 허용 → 현장(1단계)에서 확인. 기준은 '요구'이지 시공 사실이 아니므로 E0는 보수적으로 둠.")


# ------------------------------------------------------------------------------------------------ 4 공통변수
def params_sheet(wb):
    ws = wb.create_sheet("4_공통변수")
    title(ws, "공통 변수 — 값·범위·근거 (노란 칸을 바꾸면 모든 건물 시트가 다시 계산됨)",
          "몬테카를로는 하한–상한 범위(logu: 로그 균등, tri: 삼각)에서 뽑음 · fixed는 고정")
    header(ws, 4, ["키", "기호", "항목", "단위", "값", "하한", "상한", "분포", "구분", "근거"],
           widths=[10, 12, 46, 10, 11, 10, 10, 8, 9, 70], height=26)
    for k, p in enumerate(D["params"], start=5):
        vals = [p["key"], p["symbol"], p["item"], p["unit"], p["best"], p["low"], p["high"], p["distribution"], p["group"], p["basis"]]
        for c, v in enumerate(vals, start=1):
            v = None if (isinstance(v, float) and v != v) else v
            small = isinstance(v, float) and v != 0 and abs(v) < 0.001
            rc(ws, k, c, v, font=F_INPUT if c == 5 else F_BASE, fill=FILL_IN if c == 5 else None,
               fmt=(SCI if small else None) if c in (5, 6, 7) else None)
        define(wb, f"P_{p['key']}", "4_공통변수", f"$E${k}")


# ------------------------------------------------------------------------------------------------ building sheets
def building_sheet(wb, name, b, py=None):
    ws = wb.create_sheet(name)
    cols = []
    for z in b["zones"] if b else []:
        if z["col"] not in cols:
            cols.append(z["col"])
    colno = {c: i + 1 for i, c in enumerate(cols)}
    title(ws, f"{name} — {b['name'] if b else '내 건물'}",
          (b["form"] + " (가상 예제)") if b else "노란 칸에 건물·구간을 넣으면 결과가 계산됩니다 · 비어 있는 구간 행은 계산에서 빠짐")
    for c, w in zip("ABCDEFGHI", [30, 16, 30, 8, 26, 9, 15, 30, 14]):   # shared by the top panel and the zone tables
        ws.column_dimensions[c].width = w

    # building inputs
    spr = b["sprinkler"] if b else False
    spr_from = 1 if spr is True else (int(spr["from"]) if isinstance(spr, dict) else 0)
    pil = b.get("piloti") if b else None
    binp = [("건물", b["name"] if b else ""), ("형태", b["form"] if b else ""), ("층수 N (≤ 30)", b["N"] if b else None),
            ("층고 h (m)", b["h"] if b else 3.0), ("고가사다리 유효 높이 (m)", "=P_h_reach"), ("허가 연도", b["permit"] if b else None),
            ("스프링클러 시작층 (0 없음 · 1 전층 · k = k층 이상)", spr_from), ("필로티 (Y/N)", "Y" if pil else "N"),
            ("필로티 천장 가연 (Y/N)", ("Y" if pil["soffit_combustible"] else "N") if pil else "N"),
            ("1층 방화구획 (Y/N)", ("Y" if pil["ground_compartment"] else "N") if pil else "Y")]
    band(ws, 3, "건물", 3)
    for k, (lab, v) in enumerate(binp, start=4):
        rc(ws, k, 1, lab, font=F_BOLD)
        rc(ws, k, 2, v, font=F_INPUT, fill=FILL_IN)
    ws.merge_cells("B4:C4")
    ws.merge_cells("B5:C5")
    # remediation switches
    rc(ws, 3, 5, "대책 스위치 (Y/N)", font=F_SUB, fill=FILL_GREY)
    rc(ws, 3, 6, "", fill=FILL_GREY)
    for k, lab in enumerate(["개구부 차단재", "화재확산방지구조", "하단 보호·필로티 개선", "스프링클러 전층 설치", "외장재 불연 교체"], start=5):
        rc(ws, k, 5, lab, font=F_BOLD)
        rc(ws, k, 6, "N", font=F_INPUT, fill=FILL_IN, align=CENTER)
    for k, (lab, f) in enumerate([("적용 스프링클러 시작층", '=IF($F$8="Y",1,$B$10)'), ("적용 필로티 천장 가연", '=IF($F$7="Y","N",$B$12)'),
                                   ("적용 1층 방화구획", '=IF($F$7="Y","Y",$B$13)')], start=11):
        rc(ws, k, 5, lab, font=F_NOTE)
        rc(ws, k, 6, f, align=CENTER)
    # results
    rc(ws, 3, 8, "결과", font=F_SUB, fill=FILL_GREY)
    rc(ws, 3, 9, "", fill=FILL_GREY)
    zc = lambda col: f"{col}{Z_CALC + 1}:{col}{Z_CALC + NZ}"
    res = [("R (노출 인원/년)", f"=SUM({zc('Y')})+SUM({zc('AA')})+SUM({zc('AB')})", SCI),
           ("R_ref (외벽 모두 불연)", f"=SUM({zc('Z')})", SCI),
           ("재실자 수", f"=SUM({zc('W')})", "0"),
           ("IR (/인·년)", '=IFERROR(I5/I7,"")', SCI),
           ("r = IR / λ_res", '=IFERROR(I8/P_lam_res,"")', "0.0000"),
           ("등급 (1–5)", '=IF(N(I7)=0,"",1+(I9>=P_g1)+(I9>=P_g2)+(I9>=P_g3)+(I9>=P_g4))', "0"),
           ("외벽 노출 빈도 F_ign (/년)", f"=SUM({zc('X')})+SUM({zc('R')})+SUM({zc('S')})", SCI),
           ("노출 1건당 노출 인원 C̄", '=IFERROR(I5/I11,"")', "0.0"),
           ("표시", '=IF(I10="","",I10&"등급("&CHOOSE(I10,"낮음","보통","높음","매우 높음","심각")&")"&IF(I12>=P_m_flag,"·M","")'
                   f'&IF(MAX({zc("F")})=1,"·A",""))', None),
           ("외벽 위험 배수 RR = R / R_ref", '=IFERROR(I5/I6,"")', "0.0"),
           ("외벽 재료 기여율 AF", '=IFERROR(1-I6/I5,"")', "0.0%"),
           ("ΔR = R − R_ref (교체로 줄일 위험)", "=I5-I6", SCI)]
    for k, (lab, f, fmt) in enumerate(res, start=5):
        rc(ws, k, 8, lab, font=F_BOLD)
        key = k in (9, 10, 13)
        rc(ws, k, 9, f, fmt=fmt, fill=FILL_KEY if key else None, font=F_KEY if key else F_BASE)

    # zone inputs
    band(ws, Z_IN - 1, "구간 입력 (한 행 = 같은 외벽 시스템·같은 노출의 수직 구간)", 24)
    zin = ["번호", "ID", "구간 이름", "열", "시작층 b", "끝층 t", "외벽 유형(키)", "용도 res/com", "층당 세대 n_u", "층당 면적 m²",
           "노출 재실자/층", "창 계수", "층간 틈 미흡", "기존 차단재", "기존 확산방지", "하단 보호", "필로티 주차면", "기타 외부 화원",
           "TGA r_max,air", "TGA f₄₀₀,air", "콘 판정 sc/c", "시험 t₆₀₀ (s)", "시험 p_B", "근거·메모"]
    header(ws, Z_IN, zin, height=40)
    for k in range(10, 24):
        ws.column_dimensions[L(k)].width = 9
    ws.column_dimensions["X"].width = 42
    ws.row_dimensions[10].height = 30
    zones = b["zones"] if b else []
    for z in range(1, NZ + 1):
        r = Z_IN + z
        rc(ws, r, 1, z, align=CENTER)
        zz = zones[z - 1] if z <= len(zones) else None
        ev = (b.get("evidence_levels", {}).get(zz.get("evidence", ""), {}) if b else {}) if zz else {}
        tga = zz.get("tga") if zz else None
        t600 = None
        pB = None
        if zz and zz.get("evidence") == "E3" and zz["type"].startswith("CAV_"):
            t600 = None   # the type row already holds the test values
        vals = [zz["id"], zz["name"], colno[zz["col"]], zz["b"], zz["t"], zz["type"], zz["use"], zz.get("n_u", 0) or None,
                zz.get("area"), zz["occ_band"], zz.get("win", 1.0), "Y" if zz.get("gap") else "N", "Y" if zz.get("closer") else "N",
                "Y" if zz.get("barrier") else "N", "Y" if zz.get("base_closed") else "N",
                next((e["count"] for e in zz.get("ext", []) if e["src"] == "piloti"), None),
                next((e["count"] for e in zz.get("ext", []) if e["src"] == "adjacent"), None),
                tga["r_max_air"] if tga else None, tga["f400_air"] if tga else None, tga["cone"] if tga else None,
                t600, pB, f"{zz.get('evidence', '')} {zz.get('note', '')}".strip()] if zz else [None] * 23
        for c, v in enumerate(vals, start=2):
            rc(ws, r, c, v, font=F_INPUT if c < 24 else F_NOTE, fill=FILL_IN if c < 24 else None,
               align=CENTER if c not in (3, 24) else WRAP)
        ws.row_dimensions[r].height = 30

    # zone calculation
    band(ws, Z_CALC - 1, "구간 계산 (수식)", 29)
    zcalc = ["번호", "적용 유형", "외벽 경로", "t₆₀₀ (s)", "p_B", "급발 A", "차단재 인정", "확산방지 인정", "τ (분/층)", "p_e 창",
             "p_e2 하단", "p_cb", "p_L", "f_base (/층·년)", "IE1", "하부 구간", "f 하부 최상층", "F 필로티", "F 기타", "필로티 지속 하한",
             "상한", "계단 오염", "재실자", "Σ f_i", "R1", "R1_ref", "R2", "R3", "재료값 출처"]
    header(ws, Z_CALC, zcalc, height=40)
    ws.column_dimensions["AC"].width = 30
    for z in range(1, NZ + 1):
        r, i = Z_CALC + z, Z_IN + z
        g = G0 + (z - 1) * BLOCK
        act, ref, row2 = (g + 34, g + 63), (g + 67, g + 96), g + 100
        f = {
            "A": z,
            "B": f'=IF($B{i}="","",IF(AND($F$9="Y",VLOOKUP($G{i},T_types,5,FALSE)=1),"NC",$G{i}))',
            "C": f'=IF($B{r}="",0,VLOOKUP($B{r},T_types,5,FALSE))',
            "D": (f'=IF($C{r}<>1,"",IF($V{i}<>"",$V{i},IF($U{i}="c",VLOOKUP("C_UNKNOWN",T_types,6,FALSE),'
                  f'IF(AND($U{i}="sc",N($S{i})>0,N($T{i})>=P_f_melt),P_tga_c0+P_tga_c1/$S{i},VLOOKUP($B{r},T_types,6,FALSE)))))'),
            "E": (f'=IF($C{r}<>1,"",IF($W{i}<>"",$W{i},IF($U{i}="c",VLOOKUP("C_UNKNOWN",T_types,7,FALSE),'
                  f'IF(AND($U{i}="sc",N($S{i})>0,N($T{i})>=P_f_melt),MIN(0.95,MAX(0.05,($T{i}-P_f_lo)/(P_f_hi-P_f_lo))),'
                  f'VLOOKUP($B{r},T_types,7,FALSE)))))'),
            "F": (f'=IF($C{r}<>1,0,IF($U{i}="c",1,IF(AND($U{i}="sc",N($T{i})>0,N($T{i})<P_f_melt),1,'
                  f'IF(AND($U{i}="sc",N($S{i})>0,N($T{i})>=P_f_melt),0,VLOOKUP($B{r},T_types,8,FALSE)))))'),
            "G": f'=IF($C{r}=1,VLOOKUP($B{r},T_types,10,FALSE),0)',
            "H": f'=IF($C{r}=1,VLOOKUP($B{r},T_types,9,FALSE),0)',
            "I": f'=IF($C{r}=1,P_s*$D{r}/60*$B$7/3,"")',
            "J": f'=IF(AND(OR($N{i}="Y",$F$5="Y"),$G{r}=1),P_p_e1,P_p_e0)',
            "K": f'=IF(AND(OR($P{i}="Y",$F$7="Y"),$G{r}=1),P_p_e1,P_p_e0)',
            "L": f'=IF(AND(OR($O{i}="Y",$F$6="Y"),$H{r}=1),P_p_cb,0)',
            "M": f'=MIN(0.9,P_p_L*IF($M{i}="Y",P_f_gap,1))',
            "N": f'=IF($B{i}="",0,IF($H{i}="res",N($I{i})*P_lam_res,N($J{i})*P_lam_com)*P_p_fo*P_p_v*N($L{i}))',
            "O": f'=IF(AND($B{i}<>"",N($L{i})>0,OR(N($I{i})>0,N($J{i})>0)),1,0)',
            "P": (f"=SUMPRODUCT(($D${Z_IN + 1}:$D${Z_IN + NZ}=$D{i})*($F${Z_IN + 1}:$F${Z_IN + NZ}=$E{i}-1)"
                  f"*($O${Z_CALC + 1}:$O${Z_CALC + NZ}=1)*($A${Z_IN + 1}:$A${Z_IN + NZ}))"),
            "Q": (f'=IF($P{r}=0,0,INDEX($N${Z_CALC + 1}:$N${Z_CALC + NZ},$P{r})*(1-P_p_s*IF(AND($F$11>0,'
                  f'INDEX($F${Z_IN + 1}:$F${Z_IN + NZ},$P{r})>=$F$11),1,0)))'),
            "R": f'=IF(AND($B$11="Y",N($Q{i})>0),P_lam_car*$Q{i}*IF($F$12="Y",P_p_x_pc,P_p_x_pn),0)',
            "S": f"=P_lam_adj*N($R{i})*P_p_x_adj",
            "T": '=IF($F$12="Y",P_tep_lo,P_te2_lo)',
            "U": '=IF($F$12="Y",P_tep_hi,P_te2_hi)',
            "V": '=IF(AND($B$11="Y",$F$13<>"Y"),1,0)',
            "W": f'=IF($B{i}="",0,IF($H{i}="res",N($I{i})*P_occ_unit,N($J{i})*P_dens_com)*($F{i}-$E{i}+1))',
            "X": f'=IF($O{r}=1,$N{r}*(($F{i}-$E{i}+1)-P_p_s*IF($F$11>0,MAX(0,$F{i}-MAX($E{i},$F$11)+1),0)),0)',
            "Y": f"=SUMPRODUCT($AF${act[0]}:$AF${act[1]},$AG${act[0]}:$AG${act[1]})",
            "Z": f"=SUMPRODUCT($AF${ref[0]}:$AF${ref[1]},$AG${ref[0]}:$AG${ref[1]})",
            "AA": f"=$R{r}*$AG${row2}+$S{r}*$AG${row2 + 1}",
            "AB": f"=$Q{r}*$AG${row2 + 2}",
            "AC": (f'=IF($C{r}<>1,IF($B{r}="","","외벽 경로 없음"),IF($V{i}<>"","시스템 시험값(E3)",IF($U{i}="c","콘: 가연 → 가연 경계값(E2)",'
                   f'IF(AND($U{i}="sc",N($S{i})>0,N($T{i})>=P_f_melt),"TGA 예측식(E2)",VLOOKUP($B{r},T_types,2,FALSE)&" 기본값"))))'),
        }
        fmts = {"D": "0", "E": "0.000", "I": "0.000", "J": "0.00", "K": "0.00", "L": "0.00", "M": "0.000", "N": SCI, "Q": SCI,
                "R": SCI, "S": SCI, "W": "0", "X": SCI, "Y": SCI, "Z": SCI, "AA": SCI, "AB": SCI}
        for col, v in f.items():
            cell(ws, f"{col}{r}", v, fmt=fmts.get(col), align=CENTER if col != "AC" else WRAP,
                 fill=FILL_KEY if col in ("Y", "Z", "AA", "AB") else None)
        ws.row_dimensions[r].height = 28

    # storey table
    band(ws, K0 - 1, "층별 표 (k = 1–30): 고가사다리 도달, 스프링클러, 열별 노출 재실자 c(k)와 비주거 여부", 8)
    header(ws, K0, ["층 k", "바닥 높이 (m)", "사다리 도달 1/0", "스프링클러 1/0", "c(k) 열1", "비주거 열1", "c(k) 열2", "비주거 열2"], height=30)
    zi = lambda col: f"${col}${Z_IN + 1}:${col}${Z_IN + NZ}"
    for k in range(1, NS + 1):
        r = K0 + k
        rc(ws, r, 1, k, align=CENTER)
        rc(ws, r, 2, f"=(A{r}-1)*$B$7", fmt="0.0", align=CENTER)
        rc(ws, r, 3, f"=IF(B{r}<=$B$8,1,0)", align=CENTER)
        rc(ws, r, 4, f"=IF(AND($F$11>0,A{r}>=$F$11),1,0)", align=CENTER)
        for cn, (cc, cf) in enumerate([(5, 6), (7, 8)], start=1):
            rc(ws, r, cc, f"=SUMPRODUCT(({zi('D')}={cn})*({zi('E')}<=A{r})*({zi('F')}>=A{r})*{zi('K')})", fmt="0.0", align=CENTER)
            rc(ws, r, cf, f'=IF(SUMPRODUCT(({zi("D")}={cn})*({zi("E")}<=A{r})*({zi("F")}>=A{r})*({zi("H")}="com"))>0,1,0)', align=CENTER)

    # grid blocks
    for z in range(1, NZ + 1):
        zone_block(ws, z)
    ws.freeze_panes = None
    return ws


def zone_block(ws, z):
    g = G0 + (z - 1) * BLOCK
    i_, r_ = Z_IN + z, Z_CALC + z
    b, t, col = f"$E${i_}", f"$F${i_}", f"$D${i_}"
    path, tau, pB, pe, pe2, pcb, pL, ie1, lower = (f"${c}${r_}" for c in ("C", "I", "E", "J", "K", "L", "M", "O", "P"))
    stair = f"$V${r_}"
    band(ws, g, f"구간 {z} — 보조표(j별) · IE1 격자(발화층 i × 확산 층수 j) · 하단 화원·하부 구간 행", 33)
    # helper table: rows g+2 … g+30 for j = 1 … 29
    hh = ["j", "t_a = jτ", "t_l = jτ_L", "S_FS,in(t_a)", "S_FS,out(t_a)", "S_FS,in(t_l)", "S_FS,out(t_l)", "S_e1(t_a)", "S_e,필로티(t_a)",
          "S_e2(t_a)", "q_주거(t_a)", "q_비주거(t_a)", "q_주거(t_l)", "q_비주거(t_l)", "q_주거(min)", "q_비주거(min)", "(1−p_cb)^j",
          "p_L^j", "p_L,하부^j"]
    header(ws, g + 1, hh, col0=2, height=30)
    h0, h1 = g + 2, g + 1 + NJ
    for j in range(1, NJ + 1):
        r = g + 1 + j
        Sfs = lambda tt, m: f"1-NORMSDIST((LN({tt})-LN({m}))/P_sig_fs)"
        q = lambda tt, m: f"1-NORMSDIST((LN({tt})-LN({m}))/P_sig_r)"
        se = lambda tt, lo, hi: f"MAX(0,MIN(1,({hi}-{tt})/({hi}-{lo})))"
        vals = [j, f'=IF({path}=1,B{r}*{tau},"")', f"=B{r}*P_tau_L",
                f"=IF({path}=1,{Sfs(f'C{r}', 'P_t_fs_lo')},0)", f"=IF({path}=1,{Sfs(f'C{r}', 'P_t_fs_hi')},0)",
                f"={Sfs(f'D{r}', 'P_t_fs_lo')}", f"={Sfs(f'D{r}', 'P_t_fs_hi')}",
                f"=IF({path}=1,{se(f'C{r}', 'P_te1_lo', 'P_te1_hi')},0)", f"=IF({path}=1,{se(f'C{r}', f'$T${r_}', f'$U${r_}')},0)",
                f"=IF({path}=1,{se(f'C{r}', 'P_te2_lo', 'P_te2_hi')},0)",
                f"=IF({path}=1,{q(f'C{r}', 'P_rset_res')},0)", f"=IF({path}=1,{q(f'C{r}', 'P_rset_com')},0)",
                f"={q(f'D{r}', 'P_rset_res')}", f"={q(f'D{r}', 'P_rset_com')}",
                f"=IF({path}=1,{q(f'MIN(C{r},D{r})', 'P_rset_res')},0)", f"=IF({path}=1,{q(f'MIN(C{r},D{r})', 'P_rset_com')},0)",
                f"=(1-{pcb})^B{r}", f"={pL}^B{r}",
                f"=IF({lower}>0,INDEX($M${Z_CALC + 1}:$M${Z_CALC + NZ},{lower})^B{r},0)"]
        for c, v in enumerate(vals, start=2):
            rc(ws, r, c, v, font=F_SMALL, fmt="0.000" if c > 2 else None, align=CENTER)
    H = {name: f"${L(c)}${h0}:${L(c)}${h1}" for name, c in zip(
        ["ta", "tl", "fin_a", "fout_a", "fin_l", "fout_l", "se1", "sep", "se2", "qr_a", "qc_a", "qr_l", "qc_l", "qr_m", "qc_m", "cb", "pl", "ply"],
        range(3, 21))}
    ST = lambda c: f"${c}${K0 + 1}:${c}${K0 + NS}"
    ck = lambda k: f"IF({col}=1,INDEX({ST('E')},{k}),INDEX({ST('G')},{k}))"
    com = lambda k: f"IF({col}=1,INDEX({ST('F')},{k}),INDEX({ST('H')},{k}))"
    reach = lambda k: f"INDEX({ST('C')},{k})"
    sel = lambda k, jj, a, bb: f"IF({reach(k)}=1,INDEX({H[a]},{jj}),INDEX({H[bb]},{jj}))"
    selq = lambda k, jj, res, cm: f"IF({com(k)}=1,INDEX({H[cm]},{jj}),INDEX({H[res]},{jj}))"

    # IE1 grids (with façade path / reference)
    for gi, (lab, top) in enumerate([("IE1 격자 — 외벽 경로 + 창→창 도약 (값 = c·[…], 행 합 × f_i = R1)", g + 32),
                                     ("IE1 기준 격자 — 외벽 불연(창→창 도약만), 행 합 × f_i = R1_ref", g + 65)]):
        rc(ws, top, 2, lab, font=F_BOLD, border=False)
        rc(ws, top + 1, 2, "i \\ j", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        for j in range(1, NJ + 1):
            rc(ws, top + 1, 2 + j, j, font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        rc(ws, top + 1, 32, "f_i (/년)", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        rc(ws, top + 1, 33, "행 합", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
        for i in range(1, NS + 1):
            r = top + 1 + i
            rc(ws, r, 2, i, font=F_BOLD, align=CENTER)
            for j in range(1, NJ + 1):
                jj = f"{L(2 + j)}${top + 1}"
                k = f"($B{r}+{jj})"
                nspr = f"IF($F$11>0,MAX(0,{k}-MAX($B{r}+1,$F$11)+1),0)"
                bj = f"INDEX({H['pl']},{jj})*(1-P_p_s)^{nspr}*{sel(k, jj, 'fin_l', 'fout_l')}"
                qb = selq(k, jj, "qr_l", "qc_l")
                cond = f"AND({ie1}=1,$B{r}>={b},$B{r}<={t},{k}<=$B$6)"
                if gi == 1:
                    fml = f"=IF({cond},{ck(k)}*{bj}*{qb},0)"
                else:
                    aj = f"INDEX({H['cb']},{jj})*{sel(k, jj, 'fin_a', 'fout_a')}*({pB}+(1-{pB})*INDEX({H['se1']},{jj}))"
                    pa = f"({pe}*{aj})"
                    qa, qm = selq(k, jj, "qr_a", "qc_a"), selq(k, jj, "qr_m", "qc_m")
                    fml = (f"=IF({cond},{ck(k)}*IF(AND({path}=1,{k}<={t}),{pa}*(({bj})*{qm}+(1-({bj}))*{qa})+(1-{pa})*({bj})*{qb},"
                           f"({bj})*{qb}),0)")
                rc(ws, r, 2 + j, fml, font=F_SMALL, fmt="0.0E+0", align=CENTER, border=False)
            rc(ws, r, 32, f"=IF(AND({ie1}=1,$B{r}>={b},$B{r}<={t}),$N${r_}*(1-P_p_s*INDEX({ST('D')},$B{r})),0)", font=F_SMALL, fmt=SCI)
            rc(ws, r, 33, f"=SUM(C{r}:AE{r})", font=F_SMALL, fmt="0.000")

    # IE2 / IE3 rows
    top = g + 98
    rc(ws, top, 2, "하단 화원(IE2: 필로티 · 기타)과 하부 구간 분출화염(IE3) — 첫 층 k = b − 1 + j, 행 합 × F = R2 · R3", font=F_BOLD, border=False)
    rc(ws, top + 1, 2, "행 \\ j", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
    for j in range(1, NJ + 1):
        rc(ws, top + 1, 2 + j, j, font=F_HEAD, fill=FILL_HEAD, align=CENTER)
    rc(ws, top + 1, 33, "행 합", font=F_HEAD, fill=FILL_HEAD, align=CENTER)
    for n, lab in enumerate(["IE2 필로티", "IE2 기타", "IE3 하부 구간"]):
        r = top + 2 + n
        rc(ws, r, 2, lab, font=F_BOLD, align=CENTER)
        for j in range(1, NJ + 1):
            jj = f"{L(2 + j)}${top + 1}"
            k = f"({b}-1+{jj})"
            base_ok = f"AND({path}=1,{k}<={t},{k}<=$B$6)"
            if n < 2:
                se_col = "sep" if n == 0 else "se2"
                aj = f"INDEX({H['cb']},{jj})*{sel(k, jj, 'fin_a', 'fout_a')}*({pB}+(1-{pB})*INDEX({H[se_col]},{jj}))"
                qa = selq(k, jj, "qr_a", "qc_a")
                qv = f"({qa}+(1-{qa})*P_p_stair)" if n == 0 else qa
                if n == 0:
                    qv = f"IF({stair}=1,{qv},{qa})"
                F = f"$R${r_}" if n == 0 else f"$S${r_}"
                fml = f"=IF(AND({base_ok},{F}>0),{ck(k)}*{pe2}*{aj}*{qv},0)"
            else:
                nspr = f"IF($F$11>0,MAX(0,{k}-MAX({b},$F$11)+1),0)"
                bj = f"(INDEX({H['ply']},{jj})*(1-P_p_s)^{nspr}*{sel(k, jj, 'fin_l', 'fout_l')})"
                aj = f"INDEX({H['cb']},{jj})*{sel(k, jj, 'fin_a', 'fout_a')}*({pB}+(1-{pB})*INDEX({H['se1']},{jj}))"
                qa, qm, qb = selq(k, jj, "qr_a", "qc_a"), selq(k, jj, "qr_m", "qc_m"), selq(k, jj, "qr_l", "qc_l")
                fml = f"=IF(AND({base_ok},{lower}>0),{ck(k)}*{pe2}*{aj}*({bj}*{qm}+(1-{bj})*{qa}-{bj}*{qb}),0)"
            rc(ws, r, 2 + j, fml, font=F_SMALL, fmt="0.0E+0", align=CENTER, border=False)
        rc(ws, r, 33, f"=SUM(C{r}:AE{r})", font=F_SMALL, fmt="0.000")
    for c in range(3, 32):
        ws.column_dimensions[L(c)].width = max(ws.column_dimensions[L(c)].width or 0, 8)


# ------------------------------------------------------------------------------------------------ value sheets
def remediation_sheet(wb):
    ws = wb.create_sheet("9_대책비교")
    title(ws, "대책별 위험도와 등급 (ebfra.py 값) — 건물 시트의 대책 스위치(F5–F9)로 같은 값을 재현",
          "차단재·확산방지구조는 가연재가 바깥 면이거나 패널 심재로 이어진 유형(PE 심재 복합패널 등)에서 효과를 인정하지 않음")
    header(ws, 4, ["건물", "대책", "R (명/년)", "IR (/인·년)", "r", "등급·표시", "C̄ (명/건)", "ΔR (명/년)", "R 감소율"],
           widths=[8, 34, 12, 12, 10, 22, 10, 12, 10], height=28)
    base = {s["building"]: s["R"] for s in D["summary"]}
    for k, o in enumerate(D["remediation"], start=5):
        vals = [o["building"], o["option"], o["R"], o["IR"], o["r"], o["label"], o["Cbar"], o["dR"], 1 - o["R"] / base[o["building"]]]
        for c, v in enumerate(vals, start=1):
            rc(ws, k, c, v, fmt={3: SCI, 4: SCI, 5: "0.0000", 7: "0.0", 8: SCI, 9: "0%"}.get(c),
               fill=FILL_KEY if o["option"] == "현재" else None)


def evidence_sheet(wb):
    ws = wb.create_sheet("10_증거수준")
    title(ws, "증거 수준별 결과 — 예제 B(15층 아파트): 사진·현장 → 시료(TGA) → 시스템 시험",
          "점추정은 각 수준의 재료값, 등급 확률은 재료값의 수준별 불확실성 + 건물 변수의 불확실성(몬테카를로)")
    header(ws, 4, ["수준", "t₆₀₀ (s)", "p_B", "r", "등급·표시", "근거"], widths=[8, 10, 8, 10, 18, 80], height=26)
    for k, e in enumerate(D["evidence"], start=5):
        for c, v in enumerate([e["level"], e["t600"], e["pB"], e["r"], e["label"], e["note"]], start=1):
            rc(ws, k, c, v, fmt={2: "0", 3: "0.000", 4: "0.0000"}.get(c))
        ws.row_dimensions[k].height = 30
    r0 = 6 + len(D["evidence"])
    header(ws, r0, ["수준", "ln t₆₀₀ 편차", "r 5 %", "r 50 %", "r 95 %", "P(1등급)", "P(2등급)", "P(3등급)", "P(4등급)", "P(5등급)"], height=26)
    for c in range(7, 11):
        ws.column_dimensions[L(c)].width = 10
    for k, e in enumerate(D["evidence_mc"], start=r0 + 1):
        vals = [e["level"], e["sigma_ln_t600"], e["r_p05"], e["r_p50"], e["r_p95"], e["P_g1"], e["P_g2"], e["P_g3"], e["P_g4"], e["P_g5"]]
        for c, v in enumerate(vals, start=1):
            rc(ws, k, c, v, fmt="0.000" if c == 2 else ("0.0000" if c in (3, 4, 5) else ("0%" if c > 5 else None)))
    note(ws, r0 + 6, "E0·E1의 점추정은 보수값(준불연 미상 512 s), 불확실성 분석은 준불연 모집단(중앙값 약 920 s, 편차 0.506)에서 뽑음.")


def uncertainty_sheet(wb):
    ws = wb.create_sheet("11_불확실성")
    title(ws, f"불확실성과 민감도 (ebfra.py, 몬테카를로 {D['n_mc']}회)", "모든 범위 변수를 함께 뽑아 등급 확률을 냄 · 민감도는 한 변수씩 하한·상한")
    header(ws, 4, ["건물", "r 5 %", "r 50 %", "r 95 %", "P(1)", "P(2)", "P(3)", "P(4)", "P(5)", "P(M)", "R 90 % 폭(배)", "r 90 % 폭(배)", "RR 90 % 폭(배)"],
           widths=[8, 10, 10, 10, 8, 8, 8, 8, 8, 8, 12, 12, 12], height=28)
    for k, u in enumerate(D["uncertainty"], start=5):
        vals = [u["building"], u["r_p05"], u["r_p50"], u["r_p95"], u["P_g1"], u["P_g2"], u["P_g3"], u["P_g4"], u["P_g5"], u["P_M"],
                u["R_ratio90"], u["r_ratio90"], u["RR_ratio90"]]
        for c, v in enumerate(vals, start=1):
            rc(ws, k, c, v, fmt="0.0000" if c in (2, 3, 4) else ("0%" if 5 <= c <= 10 else ("0.0" if c > 10 else None)))
    r0 = 10
    header(ws, r0, ["건물", "순위", "기호", "변수", "하한", "상한", "r(하한)", "r(상한)", "폭(log₁₀)"], height=26)
    ws.column_dimensions["D"].width = 40
    k = r0 + 1
    for bid in ("A", "B", "C"):
        rows = [s for s in D["sensitivity"] if s["building"] == bid][:6]
        for n, s in enumerate(rows, start=1):
            vals = [bid, n, s["symbol"], s["item"], s["low"], s["high"], s["r_low"], s["r_high"], s["swing"]]
            for c, v in enumerate(vals, start=1):
                small = isinstance(v, float) and v != 0 and abs(v) < 0.001
                rc(ws, k, c, v, fmt=(SCI if small else ("0.0000" if c in (7, 8) else ("0.00" if c == 9 else None))))
            k += 1
    r1 = k + 1
    header(ws, r1, ["s (실물/시험)"] + [f"r · {b}" for b in "ABC"] + [f"등급 · {b}" for b in "ABC"], height=26)
    sv = sorted({s["s"] for s in D["scale"]})
    for n, s in enumerate(sv, start=r1 + 1):
        rc(ws, n, 1, s, fmt="0.000")
        for c, bid in enumerate("ABC", start=2):
            row = next(x for x in D["scale"] if x["building"] == bid and x["s"] == s)
            rc(ws, n, c, row["r"], fmt="0.0000")
            rc(ws, n, c + 3, row["grade"])
    note(ws, r1 + len(sv) + 2, "등급을 좌우하는 자료가 건물 형태마다 다름: A 필로티는 외부 화원 빈도, B 중층 준불연은 s(시뮬레이션·실대형), C 고층 가연은 스프링클러·최성기 확률.")


def verify_sheet(wb, sheets):
    ws = wb.create_sheet("12_검증")
    title(ws, "검증 — 건물 시트의 엑셀 수식과 ebfra.py 값 비교", "같은 값을 내면 계산 단계의 오류가 없다는 근거(NFPA 551 8.2.1.1)")
    header(ws, 4, ["건물", "항목", "파이썬", "엑셀", "상대차"], widths=[10, 26, 16, 16, 12], height=24)
    items = [("R (명/년)", "R", "I5"), ("R_ref", "Rref", "I6"), ("재실자", "Nocc", "I7"), ("IR", "IR", "I8"), ("r", "r", "I9"),
             ("등급", "grade", "I10"), ("C̄", "Cbar", "I12"), ("F_ign", "F_ign", "I11")]
    k = 5
    for (sheet, bid) in sheets:
        s = next(x for x in D["summary"] if x["building"] == bid)
        for lab, key, ref in items:
            rc(ws, k, 1, bid)
            rc(ws, k, 2, lab)
            rc(ws, k, 3, s[key], fmt=SCI)
            rc(ws, k, 4, f"='{sheet}'!{ref}", fmt=SCI)
            rc(ws, k, 5, f"=IFERROR(ABS(D{k}-C{k})/MAX(ABS(C{k}),1E-300),\"\")", fmt=SCI)
            k += 1
    rc(ws, k + 1, 2, "최대 상대차", font=F_BOLD)
    rc(ws, k + 1, 5, f"=MAX(E5:E{k - 1})", fmt=SCI, fill=FILL_KEY)
    rc(ws, k + 1, 6, f'=IF(E{k + 1}<1E-6,"일치","확인 필요")', fill=FILL_KEY)


def validations(wb, sheets):
    yn = '"Y,N"'
    for name in sheets:
        ws = wb[name]
        rules = [(yn, ["F5:F9", "B11:B13", f"M{Z_IN + 1}:P{Z_IN + NZ}"]), ('"res,com"', [f"H{Z_IN + 1}:H{Z_IN + NZ}"]),
                 ('"sc,c"', [f"U{Z_IN + 1}:U{Z_IN + NZ}"]), ('"1,2"', [f"D{Z_IN + 1}:D{Z_IN + NZ}"]),
                 ("T_keys", [f"G{Z_IN + 1}:G{Z_IN + NZ}"])]
        for f1, ranges in rules:
            dv = DataValidation(type="list", formula1=f1, allow_blank=True)
            ws.add_data_validation(dv)
            for rg in ranges:
                dv.add(rg)


def main():
    wb = Workbook()
    readme(wb)
    formula_sheet(wb)
    process(wb)
    types_sheet(wb)
    history_sheet(wb)
    params_sheet(wb)
    blds = {b["id"]: b for b in D["buildings"]}
    sheets = [("5_예제A", "A"), ("6_예제B", "B"), ("7_예제C", "C")]
    for name, bid in sheets:
        building_sheet(wb, name, blds[bid])
    building_sheet(wb, "8_빈양식", None)
    remediation_sheet(wb)
    evidence_sheet(wb)
    uncertainty_sheet(wb)
    verify_sheet(wb, sheets)
    validations(wb, [s for s, _ in sheets] + ["8_빈양식"])
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
