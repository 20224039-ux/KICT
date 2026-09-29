#!/usr/bin/env python3
"""중공층(PIR) 화재시험 – 화원 상부 높이별 온도이력 · 화재확산 속도 · 성장식(at², 지수, 멱법칙) 분석.

Hollow-layer (cavity) fire test: temperature logs at 2000/2500/3000 mm above the
fire source plus 600/700/800 °C arrival times at 500–3000 mm.

    python analysis.py                      # reads data/test.xlsx
    python analysis.py --xlsx other.xlsx    # same layout (sheet 'sc_2 max')

Writes figures/*.png, results/*.csv and results/summary.json next to this file.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import matplotlib

matplotlib.use("Agg")
import matplotlib.ticker  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from scipy import optimize, stats  # noqa: E402

HERE = Path(__file__).resolve().parent
SHEET = "sc_2 max"
LOG_COLUMNS = {2000: "F", 2500: "G", 3000: "H"}  # 1 Hz logs, time (s) in column B
END_OF_RECORD_S = 1800   # from 1807 s every channel falls to ~35 °C within 5-8 s (end of test) -> excluded
PHASE_I_END_S = 70       # end of the initial source-driven rise (visual, for shading only)
PLATEAU_S = (100, 340)   # phase II: quasi-steady plateau
ONSET_RISE_C = 30        # phase III starts when the running max exceeds plateau + 30 °C
THRESHOLDS = (400, 500, 600, 700, 800, 900, 1000)
SUSTAIN_S = 30           # "for at least 30 s" variant of the arrival criterion (KS F 8414 / BS 8414 style)
BURNOUT_C = 600          # burn-out front: first drop below this after the peak
PRINTED_K = {2000: 0.0012, 2500: 0.0013, 3000: 0.0014}  # exponents as printed on the Excel chart
UNIFIED_H_REF, UNIFIED_T_REF = 2000.0, 400.0

# chart tokens (dataviz reference palette, light surface)
SURFACE, INK, INK2, MUTED = "#fcfcfb", "#0b0b0b", "#52514e", "#898781"
GRID, AXIS, BAND1, BAND2 = "#e1e0d9", "#c3c2b7", "#efeee9", "#f6f5f1"
HEIGHT_COLOR = {2000: "#2a78d6", 2500: "#eb6834", 3000: "#1baf7a"}
HEIGHT_MARKER = {2000: "D", 2500: "s", 3000: "^"}
ISO_COLOR = {600: "#86b6ef", 700: "#2a78d6", 800: "#104281"}  # ordinal blue ramp (validated)
PEAK_COLOR, BURN_COLOR = "#eb6834", INK2


def rms(x):
    return float(np.sqrt(np.mean(np.square(x))))


# ----------------------------------------------------------------------------- data

def read_workbook(path: Path):
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    rows = range(2, ws.max_row + 1)
    t = np.array([ws[f"B{r}"].value for r in rows], float)
    logs = {h: np.array([ws[f"{c}{r}"].value for r in rows], float) for h, c in LOG_COLUMNS.items()}
    keep = t <= END_OF_RECORD_S
    t, logs = t[keep], {h: T[keep] for h, T in logs.items()}

    # arrival times of 600/700/800 °C for all six heights (K1:P5, temperatures in Q3:Q5)
    table = {}
    for col in "KLMNOP":
        h = int(ws[f"{col}1"].value.rsplit("_", 1)[1])
        table[h] = {int(ws[f"Q{r}"].value): float(ws[f"{col}{r}"].value) for r in (3, 4, 5)}

    # points behind the Excel exponential chart (times in K/M/O, temperatures in L/N/P, rows 14-18)
    chart = {}
    for tc, Tc in (("K", "L"), ("M", "N"), ("O", "P")):
        h = int(ws[f"{Tc}13"].value.rsplit("_", 1)[1])
        chart[h] = (np.array([ws[f"{tc}{r}"].value for r in range(14, 19)], float),
                    np.array([ws[f"{Tc}{r}"].value for r in range(14, 19)], float))
    return t, logs, table, chart


def excel_fixed_intercepts(path: Path):
    """{height: A} for exponential trendlines that use Excel's 'Set Intercept' option."""
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    out = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not re.fullmatch(r"xl/charts/chart\d+\.xml", name):
                continue
            for ser in ET.fromstring(z.read(name)).iter(f"{{{ns['c']}}}ser"):
                label, tl = ser.find("c:tx//c:v", ns), ser.find("c:trendline", ns)
                if label is None or tl is None or tl.find("c:intercept", ns) is None:
                    continue
                m = re.search(r"_(\d+)$", label.text or "")
                if m and tl.find("c:trendlineType", ns).get("val") == "exp":
                    out[int(m.group(1))] = float(tl.find("c:intercept", ns).get("val"))
    return out


# ----------------------------------------------------------------------------- metrics

def first_arrival(t, T, thr):
    idx = np.flatnonzero(T >= thr)
    return float(t[idx[0]]) if idx.size else np.nan


def sustained_arrival(t, T, thr, hold=SUSTAIN_S):
    """Start of the first period with T >= thr lasting at least `hold` seconds."""
    edges = np.diff(np.concatenate(([0], (T >= thr).astype(int), [0])))
    for s, e in zip(np.flatnonzero(edges == 1), np.flatnonzero(edges == -1)):
        if t[e - 1] - t[s] >= hold:
            return float(t[s])
    return np.nan


def smooth(y, w=11):
    return np.convolve(np.pad(y, w // 2, mode="edge"), np.ones(w) / w, mode="valid")


def phase_metrics(t, T):
    dt = float(np.median(np.diff(t)))
    in_plateau = (t >= PLATEAU_S[0]) & (t <= PLATEAU_S[1])
    plateau = float(T[in_plateau].mean())
    env = np.maximum.accumulate(T)
    i_pk = int(np.argmax(T))
    rate = np.gradient(smooth(T), t)
    after_peak = np.arange(T.size) > i_pk
    return {
        "T0": float(T[0]),
        "rate_phase_I": float(rate[t <= 80].max()),
        "plateau": plateau,
        "plateau_sd": float(T[in_plateau].std()),
        "dT_plateau": plateau - float(T[0]),
        "onset_III": float(t[np.flatnonzero((t > PLATEAU_S[1]) & (env > plateau + ONSET_RISE_C))[0]]),
        "T_peak": float(T[i_pk]),
        "t_peak": float(t[i_pk]),
        "rate_max": float(rate.max()),
        "t_rate_max": float(t[np.argmax(rate)]),
        "dur_600": float(np.sum(T >= 600) * dt),
        "dur_800": float(np.sum(T >= 800) * dt),
        "burnout": float(t[np.flatnonzero(after_peak & (T < BURNOUT_C))[0]]),
    }


def segments(arr: dict):
    """Consecutive-height velocities for one front {height: arrival time}."""
    hs = [h for h in sorted(arr) if np.isfinite(arr[h])]
    out = []
    for h1, h2 in zip(hs, hs[1:]):
        dt = arr[h2] - arr[h1]
        out.append({"from_mm": h1, "to_mm": h2, "t_from_s": arr[h1], "t_to_s": arr[h2],
                    "dt_s": dt, "v_mm_s": (h2 - h1) / dt if dt > 0 else np.inf})
    return out


# ----------------------------------------------------------------------------- growth laws

def arrival_models(t_arr, T_lev, T_amb, A_fixed, k_printed):
    """Models for the time needed to reach each threshold (thresholds fixed, time measured)."""
    L, s = np.log(T_lev), np.sqrt(T_lev - T_amb)
    k_fix = float(np.sum(t_arr * np.log(T_lev / A_fixed)) / np.sum(t_arr ** 2))  # Excel 'Set Intercept'
    r_exp, r_at2, r_lin = stats.linregress(L, t_arr), stats.linregress(s, t_arr), stats.linregress(T_lev, t_arr)
    m = {
        "excel_trendline": {"A": A_fixed, "k": k_fix, "t_hat": np.log(T_lev / A_fixed) / k_fix},
        "excel_printed": {"A": A_fixed, "k": k_printed, "t_hat": np.log(T_lev / A_fixed) / k_printed},
        "exponential": {"A": float(np.exp(-r_exp.intercept / r_exp.slope)), "k": float(1 / r_exp.slope),
                        "t_hat": r_exp.intercept + r_exp.slope * L},
        "at2": {"a": float(1 / r_at2.slope ** 2), "t0": float(r_at2.intercept),
                "t_hat": r_at2.intercept + r_at2.slope * s},
        "linear": {"rate": float(1 / r_lin.slope), "t_hat": r_lin.intercept + r_lin.slope * T_lev},
    }
    for v in m.values():
        res = t_arr - v["t_hat"]
        v["rmse_t"], v["max_abs_t"] = rms(res), float(np.max(np.abs(res)))
    return m


def power_fit(x, y):
    """Least-squares y = a x^n (x > 0), started from the log-log line."""
    ok = y > 0
    n0, lna0 = np.polyfit(np.log(x[ok]), np.log(y[ok]), 1)
    (a, n), _ = optimize.curve_fit(lambda x, a, n: a * x ** n, x, y, p0=(np.exp(lna0), n0), maxfev=20000)
    return float(a), float(n), rms(y - a * x ** n)


def growth_models(t, T, pm):
    """Phase III (onset -> peak) growth of the running-max envelope above the plateau."""
    env = np.maximum.accumulate(T)
    sel = (t > pm["onset_III"]) & (t <= pm["t_peak"])
    x, y = t[sel] - pm["onset_III"], env[sel] - pm["plateau"]
    a, n, rmse_pow = power_fit(x, y)
    a2 = float(np.sum(y * x ** 2) / np.sum(x ** 4))
    r1 = float(np.sum(y * x) / np.sum(x ** 2))
    # sensitivity: let the time origin float between ignition and the onset
    best = None
    for ts in np.arange(0.0, pm["onset_III"], 5.0):
        s2 = (t > pm["onset_III"]) & (t <= pm["t_peak"])
        try:
            fit = power_fit(t[s2] - ts, env[s2] - pm["plateau"])
        except RuntimeError:
            continue
        if best is None or fit[2] < best[1][2]:
            best = (ts, fit)
    return {"x": x, "y": y, "n": n, "a_pow": a, "rmse_pow": rmse_pow,
            "a_at2": a2, "rmse_at2": rms(y - a2 * x ** 2), "rate_lin": r1, "rmse_lin": rms(y - r1 * x),
            "free_origin_ts": float(best[0]), "free_origin_n": best[1][1], "free_origin_rmse": best[1][2]}


def unified_model(chart):
    """t = tau0 + s (h - 2000) + b ln(T/400): one heating curve shifted in time with height."""
    rows = [(h, T, tt) for h, (ts, Ts) in chart.items() for tt, T in zip(ts, Ts)]
    h, T, y = (np.array(c, float) for c in zip(*rows))
    L = np.log(T / UNIFIED_T_REF)
    dummies = np.column_stack([(h == hh).astype(float) for hh in sorted(chart)])

    def sse(X):
        beta = np.linalg.lstsq(X, y, rcond=None)[0]
        return beta, float(np.sum((y - X @ beta) ** 2))

    X3 = np.column_stack([np.ones_like(y), h - UNIFIED_H_REF, L])
    beta, sse3 = sse(X3)
    _, sse4 = sse(np.column_stack([dummies, L]))                 # common slope, free delays
    _, sse6 = sse(np.column_stack([dummies, dummies * L[:, None]]))  # separate curves per height
    n = y.size
    dof = n - 3
    se = np.sqrt(np.diag(sse3 / dof * np.linalg.inv(X3.T @ X3)))
    q = stats.t.ppf(0.975, dof)
    f_slope = ((sse4 - sse6) / 2) / (sse6 / (n - 6))
    f_linear = (sse3 - sse4) / (sse4 / (n - 4))
    tau0, s, b = beta
    return {
        "tau0_s": float(tau0), "delay_s_per_mm": float(s), "b_s": float(b), "k": float(1 / b),
        "velocity_mm_s": float(1 / s), "velocity_ci95": [float(1 / (s + q * se[1])), float(1 / (s - q * se[1]))],
        "k_ci95": [float(1 / (b + q * se[2])), float(1 / (b - q * se[2]))],
        "rmse_t": float(np.sqrt(sse3 / n)), "resid_se_t": float(np.sqrt(sse3 / dof)),
        "F_separate_k": float(f_slope), "p_separate_k": float(stats.f.sf(f_slope, 2, n - 6)),
        "F_nonlinear_delay": float(f_linear), "p_nonlinear_delay": float(stats.f.sf(f_linear, 1, n - 4)),
        "A_equiv": {int(hh): float(UNIFIED_T_REF * np.exp(-(tau0 + s * (hh - UNIFIED_H_REF)) / b)) for hh in sorted(chart)},
    }


def unified_time(u, h, T):
    return u["tau0_s"] + u["delay_s_per_mm"] * (h - UNIFIED_H_REF) + u["b_s"] * np.log(T / UNIFIED_T_REF)


# ----------------------------------------------------------------------------- figures

def style():
    plt.rcParams.update({
        "figure.facecolor": SURFACE, "axes.facecolor": SURFACE, "savefig.facecolor": SURFACE,
        "axes.edgecolor": AXIS, "axes.linewidth": 0.8, "axes.labelcolor": INK2,
        "axes.titlesize": 10.5, "axes.titleweight": "bold", "axes.titlelocation": "left",
        "axes.titlecolor": INK, "axes.labelsize": 9.5, "font.size": 9.5,
        "xtick.color": AXIS, "ytick.color": AXIS, "xtick.labelcolor": INK2, "ytick.labelcolor": INK2,
        "xtick.labelsize": 8.5, "ytick.labelsize": 8.5,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "grid.linestyle": "-",
        "axes.axisbelow": True, "axes.spines.top": False, "axes.spines.right": False,
        "legend.frameon": False, "legend.fontsize": 8.5, "legend.labelcolor": INK2,
        "lines.solid_capstyle": "round", "lines.solid_joinstyle": "round",
    })


def ring(**kw):
    """Marker with a 2px surface ring."""
    return dict(markeredgecolor=SURFACE, markeredgewidth=1.0, **kw)


def fig_histories(t, logs, pm, path):
    fig, ax = plt.subplots(figsize=(9.6, 5.4))
    ax.axvspan(0, PHASE_I_END_S, color=BAND1, lw=0, zorder=0)
    ax.axvspan(PHASE_I_END_S, 350, color=BAND2, lw=0, zorder=0)
    for thr in (400, 600, 800):
        ax.axhline(thr, color=AXIS, lw=0.7, zorder=1)
        ax.text(END_OF_RECORD_S + 8, thr, f"{thr} °C", va="center", fontsize=8, color=MUTED)
    for x, label in ((PHASE_I_END_S / 2, "I"), ((PHASE_I_END_S + 350) / 2, "II  plateau"),
                     (780, "III  spread-driven growth"), (1450, "IV–V  peak → burn-out")):
        ax.text(x, 1255, label, ha="center", fontsize=8.5, color=INK2)
    for h, T in logs.items():
        ax.plot(t, T, color=HEIGHT_COLOR[h], lw=1.4, label=f"{h} mm", zorder=3)
        m = pm[h]
        ax.plot(m["t_peak"], m["T_peak"], HEIGHT_MARKER[h], ms=7, color=HEIGHT_COLOR[h], zorder=4, **ring())
        dx, ha = (-14, "right") if h == 2000 else ((-12, "right") if h == 2500 else (12, "left"))
        ax.annotate(f"{h} mm peak {m['T_peak']:.0f} °C @ {m['t_peak']:.0f} s", (m["t_peak"], m["T_peak"]),
                    xytext=(dx, 6), textcoords="offset points", ha=ha, fontsize=8, color=INK)
    ax.set(xlim=(0, END_OF_RECORD_S), ylim=(0, 1300), xlabel="Time from ignition (s)", ylabel="Temperature (°C)")
    ax.set_title("Temperature histories above the fire source (PIR_2, sheet 'sc_2 max')")
    ax.legend(loc="upper left", bbox_to_anchor=(0.005, 0.9))
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_fronts(arrivals, pm, path):
    fig, (ax, bx) = plt.subplots(1, 2, figsize=(10.4, 5.2), sharey=True, gridspec_kw={"width_ratios": [1.6, 1]})
    fronts = [(f"{thr} °C first arrival", arrivals[thr], ISO_COLOR[thr], "o", "-") for thr in (600, 700, 800)]
    fronts += [("peak temperature", {h: pm[h]["t_peak"] for h in pm}, PEAK_COLOR, "D", "-"),
               (f"burn-out (< {BURNOUT_C} °C)", {h: pm[h]["burnout"] for h in pm}, BURN_COLOR, "s", "--")]
    for label, arr, color, mk, ls in fronts:
        hs = [h for h in sorted(arr) if np.isfinite(arr[h])]
        ax.plot([arr[h] for h in hs], hs, ls=ls, color=color, lw=1.4, marker=mk, ms=6, label=label, **ring())
        seg = [s for s in segments(arr) if np.isfinite(s["v_mm_s"])]
        bx.plot([s["v_mm_s"] for s in seg], [(s["from_mm"] + s["to_mm"]) / 2 for s in seg],
                ls=ls, color=color, lw=0.8, alpha=0.9, marker=mk, ms=6, **ring())
    ax.set(xlim=(0, 1600), ylim=(0, 3250), xlabel="Arrival time (s)", ylabel="Height above fire source (mm)")
    ax.set_yticks(range(0, 3001, 500))
    ax.set_title("(a) Front trajectories")
    ax.legend(loc="lower right")
    bx.set_xscale("log")
    bx.set(xlim=(0.8, 200), xlabel="Segment velocity Δz/Δt (mm/s, log scale)")
    bx.set_xticks([1, 2, 5, 10, 20, 50, 100])
    bx.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:g}"))
    bx.xaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    bx.text(0.98, 0.02, "segment midpoints; 500–1500 mm from the summary table",
            transform=bx.transAxes, ha="right", fontsize=7.5, color=MUTED)
    bx.axvspan(1.5, 2.5, color=BAND1, lw=0, zorder=0)
    bx.text(1.95, 3160, "1.5–2.5 mm/s", ha="center", fontsize=8, color=INK2)
    bx.set_title("(b) Velocity between adjacent levels")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_growth(t, logs, pm, gm, path):
    fig, axes = plt.subplots(1, 3, figsize=(11.4, 4.2), sharey=True)
    for ax, (h, T) in zip(axes, logs.items()):
        m, g = pm[h], gm[h]
        sel = (t > m["onset_III"]) & (t <= m["t_peak"])
        x = g["x"]
        ax.plot(t[sel] - m["onset_III"], T[sel] - m["plateau"], color=HEIGHT_COLOR[h], lw=0.8, alpha=0.35)
        ax.plot(x, g["y"], color=HEIGHT_COLOR[h], lw=1.5, label="measured (running max)")
        ax.plot(x, g["a_pow"] * x ** g["n"], color=INK, lw=1.2, label="power law  a·t'ⁿ")
        ax.plot(x, g["a_at2"] * x ** 2, color=INK, lw=1.2, ls="--", label="at'²")
        ax.plot(x, g["rate_lin"] * x, color=MUTED, lw=1.2, ls=":", label="linear")
        ax.text(0.03, 0.97, f"n = {g['n']:.2f}  (RMSE {g['rmse_pow']:.0f} °C)\n"
                            f"at'²: RMSE {g['rmse_at2']:.0f} °C\nlinear: RMSE {g['rmse_lin']:.0f} °C",
                transform=ax.transAxes, va="top", fontsize=8, color=INK)
        ax.set_title(f"{h} mm  (onset {m['onset_III']:.0f} s → peak {m['t_peak']:.0f} s)")
        ax.set_xlabel("t' = time since phase-III onset (s)")
    axes[0].set_ylabel("Rise above plateau (°C)")
    axes[0].set_ylim(0, 1000)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=4, bbox_to_anchor=(0.5, 0.0))
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_height_trends(pm, gm, decay, path):
    hs = sorted(pm)
    panels = [
        ("(a) Plateau rise ΔT (phase II)", "°C", [pm[h]["dT_plateau"] for h in hs]),
        ("(b) Max heating rate, first 80 s", "°C/s", [pm[h]["rate_phase_I"] for h in hs]),
        ("(c) Growth exponent n (phase III)", "n", [gm[h]["n"] for h in hs]),
        ("(d) Peak temperature", "°C", [pm[h]["T_peak"] for h in hs]),
        ("(e) Time of peak", "s", [pm[h]["t_peak"] for h in hs]),
        ("(f) Time spent above threshold", "s", None),
    ]
    fig, axes = plt.subplots(2, 3, figsize=(11.2, 6.4))
    for ax, (title, unit, vals) in zip(axes.flat, panels):
        if vals is None:
            for thr, key in ((600, "dur_600"), (800, "dur_800")):
                v = [pm[h][key] for h in hs]
                ax.plot(hs, v, color=ISO_COLOR[thr], lw=1.2, marker="o", ms=6, label=f"≥ {thr} °C", **ring())
                ax.annotate(f"≥ {thr} °C", (hs[-1], v[-1]), xytext=(6, 0), textcoords="offset points",
                            va="center", fontsize=8, color=INK2)
        else:
            ax.plot(hs, vals, color=AXIS, lw=1.0, zorder=2)
            for h, v in zip(hs, vals):
                ax.plot(h, v, HEIGHT_MARKER[h], ms=7, color=HEIGHT_COLOR[h], zorder=3, **ring())
                ax.annotate(f"{v:.2f}" if unit in ("n", "°C/s") else f"{v:.0f}", (h, v), xytext=(0, 8),
                            textcoords="offset points", ha="center", fontsize=8, color=INK)
        ax.set_title(title, fontsize=9.5)
        ax.set_ylabel(unit)
        ax.set_xticks(hs)
        ax.set_xlim(1800, 3250)
        ax.margins(y=0.25)
    ax = axes[0, 0]
    zz = np.linspace(1900, 3100, 50)
    ax.plot(zz, decay["dT0"] * np.exp(-(zz - 2000) / decay["lambda_mm"]), color=MUTED, lw=1.0, ls="--", zorder=1)
    ax.text(0.97, 0.9, f"ΔT ∝ exp(−z/λ), λ ≈ {decay['lambda_mm']:.0f} mm", transform=ax.transAxes,
            ha="right", fontsize=8, color=INK2)
    ax = axes[0, 2]
    for ref, lab in ((1, "linear"), (2, "t²")):
        ax.axhline(ref, color=AXIS, lw=0.8, zorder=1)
        ax.text(1820, ref, lab, va="bottom", fontsize=8, color=MUTED)
    ax.set_ylim(0, 2.3)
    for ax in axes[1]:
        ax.set_xlabel("Height above fire source (mm)")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_excel_check(chart, models, uni, path):
    fig, ax = plt.subplots(figsize=(9.6, 5.2))
    for h, (ts, Ts) in chart.items():
        c = HEIGHT_COLOR[h]
        tt = np.linspace(ts.min(), ts.max(), 200)
        ex, pr = models[h]["excel_trendline"], models[h]["excel_printed"]
        ax.plot(tt, ex["A"] * np.exp(ex["k"] * tt), color=c, lw=1.5, zorder=2)
        ax.plot(tt, pr["A"] * np.exp(pr["k"] * tt), color=c, lw=1.2, ls=":", zorder=2)
        TT = np.linspace(400, 800, 100)
        ax.plot(unified_time(uni, h, TT), TT, color=INK2, lw=1.0, ls="--", zorder=2)
        ax.plot(ts, Ts, HEIGHT_MARKER[h], ms=7, color=c, zorder=3, **ring())
        ax.annotate(f"{h} mm\nExcel: {ex['A']:.0f}·e^({ex['k']:.6f} t)\nprinted: {pr['A']:.0f}·e^({pr['k']} t)",
                    (ts[0], Ts[0]), xytext=(0, -46), textcoords="offset points", ha="left", fontsize=7.8, color=INK)
    handles = [plt.Line2D([], [], color=INK2, lw=1.5, label="Excel trendline as drawn (intercept fixed, exact k)"),
               plt.Line2D([], [], color=INK2, lw=1.2, ls=":", label="equation as printed (k rounded)"),
               plt.Line2D([], [], color=INK2, lw=1.0, ls="--", label="unified model: one curve shifted by height")]
    ax.legend(handles=handles, loc="upper left")
    ax.set(xlim=(0, 1300), ylim=(250, 900), xlabel="Arrival time (s)", ylabel="Threshold temperature (°C)")
    ax.set_title("Exponential trendlines on the 400–800 °C arrival data")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# ----------------------------------------------------------------------------- main

def write_csv(path, rows):
    rows = list(rows)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        for r in rows:
            w.writerow({k: ((round(v, 6) if np.isfinite(v) else "") if isinstance(v, float) else v)
                        for k, v in r.items()})


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, default=HERE / "data" / "test.xlsx")
    args = ap.parse_args()
    (HERE / "figures").mkdir(exist_ok=True)
    (HERE / "results").mkdir(exist_ok=True)

    t, logs, table, chart = read_workbook(args.xlsx)
    intercepts = excel_fixed_intercepts(args.xlsx)
    heights = sorted(logs)
    all_heights = sorted(table)

    # ---- phases, arrivals, velocities
    pm = {h: phase_metrics(t, logs[h]) for h in heights}
    arrivals, sustained, arrival_rows = {}, {}, []
    for thr in THRESHOLDS:
        arrivals[thr], sustained[thr] = {}, {}
        for h in all_heights:
            if h in logs:
                arrivals[thr][h] = first_arrival(t, logs[h], thr)
                sustained[thr][h] = sustained_arrival(t, logs[h], thr)
                src = "1 Hz log"
            elif thr in table[h]:
                arrivals[thr][h], sustained[thr][h], src = table[h][thr], np.nan, "summary table"
            else:
                continue
            arrival_rows.append({"threshold_C": thr, "height_mm": h, "first_arrival_s": arrivals[thr][h],
                                 f"sustained_{SUSTAIN_S}s_arrival_s": sustained[thr][h], "source": src})
    table_check = {h: {thr: table[h][thr] - arrivals[thr][h] for thr in table[h]} for h in heights}
    # sensitivity of the arrival times to an 11-s moving average (thresholds >= 400 °C)
    smoothing_shift = max(abs(first_arrival(t, smooth(logs[h]), thr) - arrivals[thr][h])
                          for h in heights for thr in THRESHOLDS
                          if np.isfinite(arrivals[thr][h]) and np.isfinite(first_arrival(t, smooth(logs[h]), thr)))
    sustained_shift = max(abs(sustained[thr][h] - arrivals[thr][h]) for h in heights for thr in THRESHOLDS
                          if np.isfinite(sustained[thr][h]) and np.isfinite(arrivals[thr][h]))

    velocity_rows = []
    fronts = {f"{thr} C": arrivals[thr] for thr in THRESHOLDS}
    fronts.update({f"{thr} C sustained {SUSTAIN_S}s": {h: sustained[thr][h] for h in heights} for thr in (500, 600, 700, 800)})
    fronts["peak"] = {h: pm[h]["t_peak"] for h in heights}
    fronts[f"burn-out <{BURNOUT_C} C"] = {h: pm[h]["burnout"] for h in heights}
    for name, arr in fronts.items():
        for s in segments(arr):
            velocity_rows.append({"front": name, **s})
    span = lambda arr, lo, hi: (hi - lo) / (arr[hi] - arr[lo])  # noqa: E731
    mean_v = {f"{thr} C": {"500-1500": span(arrivals[thr], 500, 1500), "1500-3000": span(arrivals[thr], 1500, 3000),
                           "500-3000": span(arrivals[thr], 500, 3000)} for thr in (600, 700, 800)}

    # ---- plateau decay with height
    H = np.array(heights, float)
    dTp = np.array([pm[h]["dT_plateau"] for h in heights])
    r = stats.linregress(H, np.log(dTp))
    rp = stats.linregress(np.log(H), np.log(dTp))
    decay = {"lambda_mm": float(-1 / r.slope), "dT0": float(np.exp(r.intercept + r.slope * 2000)),
             "r2": float(r.rvalue ** 2), "loss_per_500mm": float(1 - np.exp(500 * r.slope)), "power_exponent": float(rp.slope)}

    # ---- growth laws
    am = {h: arrival_models(chart[h][0], chart[h][1], pm[h]["T0"], intercepts.get(h, np.nan), PRINTED_K[h]) for h in heights}
    gm = {h: growth_models(t, logs[h], pm[h]) for h in heights}
    uni = unified_model(chart)
    for h in heights:
        ts, Ts = chart[h]
        am[h]["unified"] = {"t_hat": unified_time(uni, h, Ts)}
        am[h]["unified"]["rmse_t"] = rms(ts - am[h]["unified"]["t_hat"])
        am[h]["unified"]["max_abs_t"] = float(np.max(np.abs(ts - am[h]["unified"]["t_hat"])))
        tt = np.linspace(ts.min(), ts.max(), 400)
        ex, pr = am[h]["excel_trendline"], am[h]["excel_printed"]
        am[h]["printed_vs_drawn_max_C"] = float(np.max(np.abs(pr["A"] * np.exp(pr["k"] * tt) - ex["A"] * np.exp(ex["k"] * tt))))
    common_A = 250.0
    k_common_A = {h: float(np.sum(chart[h][0] * np.log(chart[h][1] / common_A)) / np.sum(chart[h][0] ** 2)) for h in heights}

    # ---- outputs
    write_csv(HERE / "results" / "arrival_times.csv", arrival_rows)
    write_csv(HERE / "results" / "spread_velocity.csv", velocity_rows)
    write_csv(HERE / "results" / "phase_metrics.csv", ({"height_mm": h, **pm[h]} for h in heights))
    growth_rows = []
    for h in heights:
        for name, v in am[h].items():
            if isinstance(v, dict):
                growth_rows.append({"height_mm": h, "data": "400-800 C arrival times", "model": name,
                                    "params": "; ".join(f"{k}={v[k]:.6g}" for k in ("A", "k", "a", "t0", "rate") if k in v),
                                    "rmse": v["rmse_t"], "rmse_unit": "s"})
        g = gm[h]
        for name, par, err in (("power law", f"a={g['a_pow']:.6g}; n={g['n']:.4f}", g["rmse_pow"]),
                               ("at2", f"a={g['a_at2']:.6g}", g["rmse_at2"]),
                               ("linear", f"rate={g['rate_lin']:.6g}", g["rmse_lin"]),
                               ("power law, free origin", f"ts={g['free_origin_ts']:.0f}; n={g['free_origin_n']:.4f}", g["free_origin_rmse"])):
            growth_rows.append({"height_mm": h, "data": "phase III envelope", "model": name, "params": par,
                                "rmse": err, "rmse_unit": "C"})
    write_csv(HERE / "results" / "growth_fits.csv", growth_rows)

    summary = {
        "phase_metrics": {h: pm[h] for h in heights},
        "arrival_first_s": {thr: arrivals[thr] for thr in THRESHOLDS},
        "arrival_sustained_s": {thr: sustained[thr] for thr in THRESHOLDS},
        "table_minus_log_s": table_check,
        "max_shift_smoothing_11s_s": float(smoothing_shift),
        "max_shift_sustained_30s_s": float(sustained_shift),
        "mean_velocity_mm_s": mean_v,
        "plateau_decay": decay,
        "excel_intercepts": intercepts,
        "arrival_models": {h: {k: ({kk: vv for kk, vv in v.items() if kk != "t_hat"} if isinstance(v, dict) else v)
                               for k, v in am[h].items()} for h in heights},
        "k_with_common_A_250": k_common_A,
        "phase_III_growth": {h: {k: v for k, v in gm[h].items() if k not in ("x", "y")} for h in heights},
        "unified_model": uni,
    }
    def clean(o):  # strict JSON: numpy scalars -> float, NaN -> null
        if isinstance(o, dict):
            return {k: clean(v) for k, v in o.items()}
        if isinstance(o, (list, tuple)):
            return [clean(v) for v in o]
        if isinstance(o, (float, np.floating)):
            return float(o) if np.isfinite(o) else None
        return o

    (HERE / "results" / "summary.json").write_text(
        json.dumps(clean(summary), indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")

    style()
    fig_histories(t, logs, pm, HERE / "figures" / "fig1_temperature_histories.png")
    fig_fronts(arrivals, pm, HERE / "figures" / "fig2_fire_front_velocity.png")
    fig_growth(t, logs, pm, gm, HERE / "figures" / "fig3_growth_models.png")
    fig_height_trends(pm, gm, decay, HERE / "figures" / "fig4_height_trends.png")
    fig_excel_check(chart, am, uni, HERE / "figures" / "fig5_excel_trendline_check.png")

    # ---- console summary
    print("phase metrics")
    for h in heights:
        m = pm[h]
        print(f"  {h} mm: plateau {m['plateau']:.0f}±{m['plateau_sd']:.0f} °C, phase-I max rate {m['rate_phase_I']:.2f} °C/s, "
              f"onset III {m['onset_III']:.0f} s, peak {m['T_peak']:.0f} °C @ {m['t_peak']:.0f} s, "
              f"≥600 °C {m['dur_600']:.0f} s, ≥800 °C {m['dur_800']:.0f} s, burn-out <{BURNOUT_C} °C @ {m['burnout']:.0f} s")
    print(f"plateau decay: λ = {decay['lambda_mm']:.0f} mm (R² {decay['r2']:.4f}), "
          f"-{100 * decay['loss_per_500mm']:.0f}% per 500 mm, power exponent {decay['power_exponent']:.2f}")
    print("arrival times (first / sustained 30 s), s")
    for thr in THRESHOLDS:
        print(f"  {thr:5d} °C: " + "  ".join(f"{h}:{arrivals[thr][h]:6.0f}/{sustained[thr].get(h, np.nan):5.0f}" for h in arrivals[thr]))
    print("segment velocities (mm/s)")
    for name, arr in fronts.items():
        print(f"  {name:22s} " + "  ".join(f"{s['from_mm']}-{s['to_mm']}:{s['v_mm_s']:6.2f}" for s in segments(arr)))
    print("span-averaged velocities (mm/s):", {k: {kk: round(vv, 2) for kk, vv in v.items()} for k, v in mean_v.items()})
    print("arrival-data models (RMSE in s)")
    for h in heights:
        print(f"  {h} mm: " + "  ".join(f"{k}={v['rmse_t']:.1f}" for k, v in am[h].items() if isinstance(v, dict)),
              f"| Excel k={am[h]['excel_trendline']['k']:.6f} (printed {PRINTED_K[h]}), max gap {am[h]['printed_vs_drawn_max_C']:.1f} °C"
              f" | exp k={am[h]['exponential']['k']:.6f} A={am[h]['exponential']['A']:.0f}"
              f" | at² a={am[h]['at2']['a']:.3e} t0={am[h]['at2']['t0']:.0f} s")
    print("k with a common intercept A=250:", {h: round(v, 6) for h, v in k_common_A.items()})
    print("phase-III growth (envelope above plateau)")
    for h in heights:
        g = gm[h]
        print(f"  {h} mm: n={g['n']:.2f} (RMSE {g['rmse_pow']:.1f}), at² a={g['a_at2']:.3e} (RMSE {g['rmse_at2']:.1f}), "
              f"linear {g['rate_lin']:.3f} °C/s (RMSE {g['rmse_lin']:.1f}); free origin ts={g['free_origin_ts']:.0f} s → n={g['free_origin_n']:.2f}")
    print("unified model:", {k: (round(v, 6) if isinstance(v, float) else v) for k, v in uni.items()})
    print("summary-table minus log arrival (s):", table_check)
    print(f"max arrival shift, >=400 °C: 11-s moving average {smoothing_shift:.0f} s, sustained {SUSTAIN_S} s {sustained_shift:.0f} s")


if __name__ == "__main__":
    main()
